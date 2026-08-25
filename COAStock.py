from tkinter import filedialog, messagebox
import tkinter as tk
import re
import unicodedata
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Border, PatternFill, Side
from openpyxl.utils import get_column_letter

from AppConstants import COA_STOCK_CONFIG
from COA import apply_part_mapping, fetch_orders_from_api
from CommonUtils import (
    create_progress_window,
    normalize_part_no,
    normalize_process_token,
    select_excel_save_file,
)


REPORT_COLUMNS = [
    "수주번호", "날짜", "납기", "출고기준일", "중요표시", "출고(고객)사",
    "품번", "품명", "규격",
    "수량", "할당수량", "잔여량", "재고", "비고",
]
STOCK_HISTORY_COLUMNS = [
    "수주번호", "비교품번", "품번", "납기", "출고(고객)사", "비고",
    "최초재고", "차감전재고", "차감량(잔여량)", "차감후재고", "재고충당여부",
]
AVAILABLE_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
TODAY_DUE_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DEFAULT_MAPPING_FILE = Path(
    r"\\Data-server\01. 공유문서\08.생산계획\생산계획\MES 생산계획 업로딩용"
    r"\프로그램\4. 품번매핑.xlsx"
)


def _read_shared_strings(book: zipfile.ZipFile) -> list[str]:
    path = "xl/sharedStrings.xml"
    if path not in book.namelist():
        return []

    strings = []
    with book.open(path) as source:
        for event, elem in ET.iterparse(source, events=("end",)):
            if elem.tag == f"{{{MAIN_NS}}}si":
                strings.append("".join(node.text or "" for node in elem.iter(f"{{{MAIN_NS}}}t")))
                elem.clear()
    return strings


def _worksheet_paths(book: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(book.read("xl/workbook.xml"))
    relationships = ET.fromstring(book.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in relationships.findall(f"{{{PKG_REL_NS}}}Relationship")
    }

    result = {}
    for sheet in workbook.findall(f".//{{{MAIN_NS}}}sheet"):
        rel_id = sheet.attrib.get(f"{{{DOC_REL_NS}}}id")
        target = targets.get(rel_id, "").replace("\\", "/")
        if target.startswith("/"):
            target = target.lstrip("/")
        elif not target.startswith("xl/"):
            target = f"xl/{target}"
        result[sheet.attrib["name"]] = target
    return result


def _cell_value(cell, shared_strings):
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(f"{{{MAIN_NS}}}t"))

    value_node = cell.find(f"{{{MAIN_NS}}}v")
    if value_node is None or value_node.text is None:
        return None
    if cell_type == "s":
        try:
            return shared_strings[int(value_node.text)]
        except (ValueError, IndexError):
            return None
    return value_node.text


def extract_stock_file_fast(plan_file: str, progress=None) -> pd.DataFrame:
    """openpyxl 없이 XLSX에서 설정된 품번/재고 두 열만 스트리밍한다."""
    path = Path(plan_file)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise RuntimeError("빠른 재고 추출은 .xlsx 또는 .xlsm 생산계획 파일만 지원합니다.")

    sheet_names = COA_STOCK_CONFIG.get("sheet_names", ["완성공정(실적)"])
    part_col = str(COA_STOCK_CONFIG.get("part_col", "F")).upper()
    stock_col = str(COA_STOCK_CONFIG.get("stock_col", "AD")).upper()
    start_row = int(COA_STOCK_CONFIG.get("data_start_row", 35))
    cell_ref_pattern = re.compile(r"([A-Z]+)\d+")
    records = []

    if progress:
        progress(35, "생산계획에서 품번/재고 열만 빠르게 읽는 중...")

    with zipfile.ZipFile(path) as book:
        shared_strings = _read_shared_strings(book)
        worksheet_paths = _worksheet_paths(book)
        found_sheet = False

        for sheet_name in sheet_names:
            sheet_path = worksheet_paths.get(sheet_name)
            if not sheet_path:
                continue
            found_sheet = True
            with book.open(sheet_path) as source:
                for event, row in ET.iterparse(source, events=("end",)):
                    if row.tag != f"{{{MAIN_NS}}}row":
                        continue
                    row_number = int(row.attrib.get("r", "0"))
                    if row_number >= start_row:
                        part_value = None
                        stock_value = None
                        for cell in row.findall(f"{{{MAIN_NS}}}c"):
                            match = cell_ref_pattern.match(cell.attrib.get("r", ""))
                            column = match.group(1) if match else ""
                            if column == part_col:
                                part_value = _cell_value(cell, shared_strings)
                            elif column == stock_col:
                                stock_value = _cell_value(cell, shared_strings)

                        part_no = normalize_part_no(part_value)
                        if part_no:
                            stock_number = pd.to_numeric(
                                str(stock_value or "").replace(",", ""), errors="coerce"
                            )
                            records.append({
                                "품번": part_no,
                                "재고": None if pd.isna(stock_number) else float(stock_number),
                            })
                    row.clear()

        if not found_sheet:
            raise RuntimeError(f"생산계획에서 재고 시트를 찾지 못했습니다: {sheet_names}")

    if progress:
        progress(90, f"재고 추출 완료: {len(records)}건")
    return pd.DataFrame(records, columns=["품번", "재고"])


def parse_lcs_date(remark, fallback_year) -> pd.Timestamp:
    """비고의 'LCS 날짜' 뒤에서 연/월/일 또는 월/일 형식을 읽는다."""
    text = str(remark or "")
    pattern = re.compile(
        r"LCS\s*(?:날짜)?\s*[:：]?\s*"
        r"(?:(\d{2,4})\s*[./\-년]\s*)?"
        r"(\d{1,2})\s*[./\-월]\s*"
        r"(\d{1,2})\s*일?",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    if not match:
        return pd.NaT

    year_text, month_text, day_text = match.groups()
    year = int(year_text) if year_text else int(fallback_year)
    if year < 100:
        year += 2000
    try:
        return pd.Timestamp(year=year, month=int(month_text), day=int(day_text)).normalize()
    except ValueError:
        return pd.NaT


def add_overseas_shipping_priority(report: pd.DataFrame) -> pd.DataFrame:
    """클라크 해외수주에 출고기준일과 중요 표시를 추가한다."""
    report = report.copy()
    report["출고기준일"] = pd.NaT
    report["중요표시"] = ""

    for idx, row in report.iterrows():
        customer = normalize_process_token(row.get("출고(고객)사", ""))
        if not customer.startswith("㈜클라크"):
            continue

        due_date = pd.to_datetime(row.get("납기"), errors="coerce")
        if pd.isna(due_date):
            continue

        lcs_date = parse_lcs_date(row.get("비고", ""), due_date.year)
        if pd.notna(lcs_date):
            report.at[idx, "출고기준일"] = lcs_date - pd.Timedelta(days=1)
            report.at[idx, "중요표시"] = "긴급: LCS 날짜 전일 출고"
        else:
            report.at[idx, "출고기준일"] = due_date.normalize() - pd.Timedelta(days=1)
            report.at[idx, "중요표시"] = "중요: 해외수주 납기 전일 출고"

    return report


def build_stock_by_part(plan_df: pd.DataFrame) -> pd.DataFrame:
    """생산계획에서 비교품번별 현재고를 만든다."""
    if plan_df.empty or "재고" not in plan_df.columns:
        return pd.DataFrame(columns=["비교품번", "재고"])

    stock_df = plan_df[["품번", "재고"]].copy()
    stock_df["비교품번"] = stock_df["품번"].apply(normalize_part_no)
    stock_df["재고"] = pd.to_numeric(stock_df["재고"], errors="coerce")
    stock_df = stock_df[(stock_df["비교품번"] != "") & stock_df["재고"].notna()].copy()
    if stock_df.empty:
        return pd.DataFrame(columns=["비교품번", "재고"])

    # 날짜별 계획 레코드에 같은 현재고가 반복되므로 합계 대신 최댓값을 사용한다.
    return stock_df.groupby("비교품번", as_index=False)["재고"].max()


def build_order_balance_report(
    order_df: pd.DataFrame,
    plan_df: pd.DataFrame,
    return_history: bool = False,
):
    """같은 비교품번의 수주를 납기순으로 정렬해 현재고를 순차 차감한다."""
    report = order_df.copy()
    stock_df = build_stock_by_part(plan_df)
    initial_stock = dict(zip(stock_df["비교품번"], stock_df["재고"]))
    remaining_stock = {part_no: float(stock) for part_no, stock in initial_stock.items()}
    report = report.sort_values(
        ["비교품번", "납기", "수주번호"], kind="stable", na_position="last"
    ).reset_index(drop=True)

    history_records = []
    available_stock_values = []
    for _, row in report.iterrows():
        comparison_part = normalize_part_no(row.get("비교품번"))
        starting_stock = float(initial_stock.get(comparison_part, 0))
        stock_before = float(remaining_stock.get(comparison_part, starting_stock))
        deduction = pd.to_numeric(row.get("잔여량"), errors="coerce")
        deduction = 0.0 if pd.isna(deduction) else float(deduction)
        stock_after = stock_before - deduction
        remaining_stock[comparison_part] = stock_after
        available_stock_values.append(stock_before)
        history_records.append({
            "수주번호": row.get("수주번호", ""),
            "비교품번": comparison_part,
            "품번": row.get("수주품번", row.get("품번", "")),
            "납기": row.get("납기"),
            "출고(고객)사": row.get("출고(고객)사", ""),
            "비고": row.get("비고", ""),
            "최초재고": starting_stock,
            "차감전재고": stock_before,
            "차감량(잔여량)": deduction,
            "차감후재고": stock_after,
            "재고충당여부": "가능" if stock_before >= deduction else "부족",
        })

    # 각 행의 재고는 해당 수주를 차감하기 직전의 가용 재고이다.
    report["재고"] = available_stock_values

    # 결과에는 COA 원래 품번을 표시하고 비교에만 매핑 품번을 사용한다.
    report["품번"] = report["수주품번"]
    report = add_overseas_shipping_priority(report)
    report["_정렬기준일"] = pd.to_datetime(
        report["출고기준일"], errors="coerce"
    ).fillna(pd.to_datetime(report["납기"], errors="coerce"))
    report = report.sort_values(
        ["_정렬기준일", "납기", "수주번호", "품번"],
        kind="stable",
        na_position="last",
    )
    report = report[REPORT_COLUMNS].copy()
    report = report.reset_index(drop=True)
    history_df = pd.DataFrame(history_records, columns=STOCK_HISTORY_COLUMNS)
    return (report, history_df) if return_history else report


def split_shipping_reports(report_df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """출고사와 뉴라텍 비고를 기준으로 출고 구분별 시트 데이터를 만든다."""
    normalized_customers = report_df["출고(고객)사"].apply(normalize_process_token)
    newratec_mask = normalized_customers.str.startswith("㈜뉴라텍")
    remarks = report_df["비고"].fillna("").astype(str)
    mirae_mask = newratec_mask & remarks.str.contains("미래전기", case=False)
    gpc_mask = (
        newratec_mask
        & remarks.str.contains("지피씨", case=False)
        & ~mirae_mask
    )
    hd_ansan_mask = normalized_customers.apply(
        lambda value: "HD건설기계" in value and "안산" in value
    )
    clark_mask = normalized_customers.str.startswith("㈜클라크")
    classified_mask = mirae_mask | gpc_mask | hd_ansan_mask | clark_mask
    return {
        "미래전기": report_df[mirae_mask].copy(),
        "지피씨": report_df[gpc_mask].copy(),
        "HD건설기계㈜ 안산": report_df[hd_ansan_mask].copy(),
        "클라크": report_df[clark_mask].copy(),
        "그외": report_df[~classified_mask].copy(),
    }


def _display_width(value) -> int:
    """한글 등 전각 문자를 고려한 엑셀 표시 너비를 계산한다."""
    if value is None:
        return 0
    return max(
        sum(2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1 for char in line)
        for line in str(value).splitlines() or [""]
    )


def apply_table_borders_and_autofit(ws) -> None:
    """사용 영역을 가운데 정렬/테두리 처리하고 내용에 맞춰 열 너비를 조정한다."""
    column_widths = {}
    for column_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        column_letter = get_column_letter(column_cells[0].column)
        content_widths = []
        for cell in column_cells:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value not in (None, ""):
                if cell.number_format == "yyyy-mm-dd":
                    content_widths.append(10)
                else:
                    content_widths.append(_display_width(cell.value))
        # 다른 열의 길이는 섞지 않고 현재 열의 가장 긴 내용만 사용한다.
        column_widths[column_letter] = min(max(content_widths, default=1) + 2, 255)

    for column_letter, width in column_widths.items():
        ws.column_dimensions[column_letter].width = width


def apply_table_borders_and_alignment(ws) -> None:
    """자동너비 변경 없이 사용 영역 전체에 가운데 정렬과 테두리를 적용한다."""
    for row in ws.iter_rows(
        min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")


def format_order_balance_sheet(ws, report_df: pd.DataFrame, today=None) -> int:
    """재고 충당 가능 행과 오늘 납기 행을 색상으로 강조한다."""
    today = pd.Timestamp(today).normalize() if today is not None else pd.Timestamp.today().normalize()
    highlighted_rows = set()

    header_columns = {cell.value: cell.column for cell in ws[1]}
    for column_name in ("날짜", "납기", "출고기준일"):
        column = header_columns.get(column_name)
        if column:
            for row_number in range(2, ws.max_row + 1):
                ws.cell(row_number, column).number_format = "yyyy-mm-dd"

    for row_idx, row in report_df.reset_index(drop=True).iterrows():
        due_date = pd.to_datetime(row["납기"], errors="coerce")
        balance = pd.to_numeric(row["잔여량"], errors="coerce")
        stock = pd.to_numeric(row["재고"], errors="coerce")
        if pd.notna(balance) and pd.notna(stock) and balance <= stock:
            for cell in ws[row_idx + 2]:
                cell.fill = AVAILABLE_FILL
            highlighted_rows.add(row_idx)

        if pd.notna(due_date) and due_date.normalize() == today:
            for cell in ws[row_idx + 2]:
                cell.fill = TODAY_DUE_FILL
            highlighted_rows.add(row_idx)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    apply_table_borders_and_autofit(ws)
    return len(highlighted_rows)


def format_stock_history_sheet(ws) -> None:
    """재고 차감 이력 시트에 필터와 날짜 형식을 적용한다."""
    header_columns = {cell.value: cell.column for cell in ws[1]}
    due_column = header_columns.get("납기")
    if due_column:
        for row_number in range(2, ws.max_row + 1):
            ws.cell(row_number, due_column).number_format = "yyyy-mm-dd"
    availability_column = header_columns.get("재고충당여부")
    if availability_column:
        for row_number in range(2, ws.max_row + 1):
            if ws.cell(row_number, availability_column).value == "가능":
                for cell in ws[row_number]:
                    cell.fill = AVAILABLE_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    apply_table_borders_and_alignment(ws)


def apply_excel_autofit(file_path: str) -> bool:
    """Excel을 백그라운드로 실행해 결과 시트에 실제 열 자동 맞춤을 적용한다."""
    excel = None
    workbook = None
    try:
        import win32com.client

        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(str(Path(file_path).resolve()))
        for worksheet in workbook.Worksheets:
            if worksheet.Name != "재고차감이력":
                worksheet.UsedRange.Columns.AutoFit()
        workbook.Save()
        return True
    except Exception:
        # Excel이 없거나 자동화할 수 없는 환경에서는 계산식 너비를 그대로 사용한다.
        return False
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def main():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    plan_file = filedialog.askopenfilename(
        title="생산계획 엑셀 선택 (재고 추출용)",
        filetypes=[("Excel files", "*.xlsx *.xlsm")],
    )
    if not plan_file:
        messagebox.showinfo("알림", "생산계획 파일을 선택하지 않았습니다.")
        return

    if DEFAULT_MAPPING_FILE.is_file():
        mapping_file = str(DEFAULT_MAPPING_FILE)
        mapping_source = "공용서버 자동 연결"
    else:
        messagebox.showinfo(
            "알림",
            "공용서버 품번매핑 파일에 접근할 수 없습니다.\n"
            "대신 사용할 품번매핑 파일을 선택해 주세요.",
        )
        mapping_file = filedialog.askopenfilename(
            title="품번매핑 파일 선택",
            filetypes=[("Excel files", "*.xlsx *.xlsm *.xls")],
        )
        if mapping_file:
            mapping_source = "직접 선택"
        else:
            messagebox.showinfo("알림", "품번매핑 파일을 선택하지 않아 매핑 없이 진행합니다.")
            mapping_file = None
            mapping_source = "사용 안 함"

    save_file = select_excel_save_file("결과 파일 저장", plan_file, "수주잔량_재고현황", root)
    if not save_file:
        return

    progress_win, progress = create_progress_window(root, "수주잔량과 재고 비교")
    try:
        progress(10, "서버에서 수주 데이터 받는 중...")
        order_df = fetch_orders_from_api()

        progress(20, "품번 매핑 적용 중...")
        order_df = apply_part_mapping(order_df, mapping_file)
        mapped_count = int(order_df["품번매핑여부"].eq("Y").sum())

        progress(30, f"생산계획 재고 읽는 중... (매핑 {mapped_count}건)")
        plan_df = extract_stock_file_fast(plan_file, progress=progress)

        progress(92, "수주잔량과 재고 연결 중...")
        report_df, stock_history_df = build_order_balance_report(
            order_df, plan_df, return_history=True
        )
        shipping_reports = split_shipping_reports(report_df)

        progress(96, "결과 저장 중...")
        with pd.ExcelWriter(save_file, engine="openpyxl") as writer:
            report_df.to_excel(writer, sheet_name="전체", index=False)
            highlighted_count = format_order_balance_sheet(
                writer.book["전체"], report_df
            )
            for sheet_name, sheet_df in shipping_reports.items():
                export_df = sheet_df
                if sheet_name != "클라크":
                    export_df = sheet_df.drop(
                        columns=["출고기준일", "중요표시"], errors="ignore"
                    )
                export_df.to_excel(writer, sheet_name=sheet_name, index=False)
                format_order_balance_sheet(writer.book[sheet_name], sheet_df)
            stock_history_df.to_excel(writer, sheet_name="재고차감이력", index=False)
            format_stock_history_sheet(writer.book["재고차감이력"])

        progress(98, "엑셀 열 너비 자동 맞춤 중...")
        autofit_applied = apply_excel_autofit(save_file)

        progress(100, "완료")
        progress_win.destroy()
        messagebox.showinfo(
            "알림",
            f"결과 저장 완료: {save_file}\n"
            f"품번매핑: {mapped_count}건 ({mapping_source})\n"
            f"색칠된 행: {highlighted_count}건\n"
            f"자동너비: {'Excel AutoFit 적용' if autofit_applied else '계산 너비 적용'}",
        )
    except Exception as exc:
        progress_win.destroy()
        messagebox.showerror("오류", str(exc))


if __name__ == "__main__":
    main()
