from pathlib import Path
import sys
from tkinter import Tk, messagebox
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from CommonUtils import (
    apply_quantities_by_part_left_to_right,
    find_column_by_keywords,
    normalize_part_no,
    normalize_process_token,
    select_date_range,
    select_excel_file,
    select_excel_save_file,
)
from ExtractPlan import extract_plan_sheet, load_calamine_worksheets
from AppConstants import get_plan_extraction_config


PLAN_SHEETS = (
    ("완성공정(실적)", "완성", "완성시트"),
    ("TANK공정(실적)", "TANK", "TANK공정시트"),
    ("CORE공정(실적)", "CORE", "CORE공정시트"),
)


def adjust_row(values):
    """첫 열의 음수 미달(선생산)을 뒤쪽 계획에서 날짜순으로 차감한다."""
    result = [float(v) for v in values]
    stock = 0.0
    if result and result[0] < 0:
        stock = abs(result[0])
        result[0] = 0.0

    for i in range(1, len(result)):
        if stock <= 0:
            break
        used = min(max(result[i], 0), stock)
        result[i] -= used
        stock -= used
    return result


def extract_mrp_source(file_path, sheet_name, default_process, **config_overrides):
    worksheets = load_calamine_worksheets(file_path, [sheet_name])
    if sheet_name not in worksheets:
        raise RuntimeError(f"{Path(file_path).name}: {sheet_name} 시트가 없습니다.")
    config = {
        "default_process": default_process,
        **get_plan_extraction_config("mrp_ver2", sheet_name),
        **config_overrides,
    }
    return extract_plan_sheet(worksheets[sheet_name], config)


def load_actuals_by_process(actual_file):
    df = pd.read_excel(actual_file, engine="calamine")
    part_col = find_column_by_keywords(df, ("품번",))
    qty_col = find_column_by_keywords(df, ("실적", "수량"))
    process_col = find_column_by_keywords(df, ("공정",))
    if part_col is None or qty_col is None:
        raise RuntimeError("output.xlsx에 품번/실적수량 열이 없습니다.")

    work = df[[part_col, qty_col] + ([process_col] if process_col else [])].copy()
    work[part_col] = work[part_col].apply(normalize_part_no)
    work[qty_col] = pd.to_numeric(work[qty_col], errors="coerce").fillna(0)
    if process_col:
        work["__MRP공정"] = work[process_col].apply(classify_process)
    else:
        work["__MRP공정"] = "전체"
    work = work[work[part_col] != ""]

    grouped = {}
    for process, group in work.groupby("__MRP공정"):
        grouped[process] = group.groupby(part_col)[qty_col].sum().to_dict()
    return grouped


def classify_process(value):
    text = normalize_process_token(value)
    if "CORE" in text:
        return "CORE"
    if "TANK" in text or "용접" in text or "CLINCHING" in text:
        return "TANK"
    if "완성" in text:
        return "완성"
    return "기타"


def build_mrp_sheet(extracted, start_date, end_date, actuals):
    metadata = ["품번", "모델", "고객사품번"]
    if extracted.empty:
        return pd.DataFrame(columns=metadata + ["합계"])

    df = extracted.copy()
    df["날짜"] = pd.to_datetime(df["날짜"], errors="coerce").dt.normalize()
    df["수량"] = pd.to_numeric(df["수량"], errors="coerce").fillna(0)
    df["품번"] = df["품번"].apply(normalize_part_no)

    plan = df[df["구분"] == "계획"].copy()
    if start_date is not None:
        plan = plan[plan["날짜"] >= pd.Timestamp(start_date)]
    if end_date is not None:
        plan = plan[plan["날짜"] <= pd.Timestamp(end_date)]

    plan_dates = sorted(plan["날짜"].dropna().unique())
    if not plan_dates:
        return pd.DataFrame(columns=metadata + ["합계"])

    # 기준일(기간 시작일)의 미달을 출력의 첫 날짜 열에 포함한다.
    shortage = df[df["구분"] == "미달"].copy()
    baseline = pd.Timestamp(start_date) if start_date is not None else pd.Timestamp(plan_dates[0])
    shortage = shortage[shortage["날짜"] == baseline]
    if not shortage.empty:
        shortage["날짜"] = pd.Timestamp(plan_dates[0])
        plan = pd.concat([plan, shortage], ignore_index=True)

    def first_nonempty(series):
        values = series.dropna().astype(str).str.strip()
        values = values[values != ""]
        return values.iloc[0] if not values.empty else ""

    info = (
        df.groupby("품번", as_index=False, sort=False)
        .agg({"모델": first_nonempty, "고객사품번": first_nonempty})
    )
    wide = plan.pivot_table(
        index="품번", columns="날짜", values="수량", aggfunc="sum", fill_value=0,
        sort=False,
    ).reset_index()
    result = info.merge(wide, on="품번", how="inner")
    date_cols = sorted(c for c in result.columns if isinstance(c, pd.Timestamp))
    result = result[metadata + date_cols]

    for idx in result.index:
        result.loc[idx, date_cols] = adjust_row(result.loc[idx, date_cols].tolist())

    result = apply_quantities_by_part_left_to_right(
        result,
        actuals,
        part_col="품번",
        value_start_idx=len(metadata),
    )
    date_cols = [c for c in result.columns if isinstance(c, pd.Timestamp)]
    result["합계"] = result[date_cols].sum(axis=1)
    result = result[result["합계"] != 0].reset_index(drop=True)
    result = result.rename(columns={c: c.strftime("%Y/%m/%d") for c in date_cols})

    total = {col: "" for col in result.columns}
    total["품번"] = ""
    for col in [c.strftime("%Y/%m/%d") for c in date_cols] + ["합계"]:
        total[col] = result[col].sum()
    return pd.concat([result, pd.DataFrame([total])], ignore_index=True)


def build_calculation_trace(extracted, start_date, end_date, actuals, output_sheet):
    """공정별 계산 과정을 품번·날짜별로 펼쳐서 보여준다."""
    columns = [
        "결과시트", "품번", "모델", "고객사품번", "날짜",
        "계획수량", "기준일미달", "미달합산후수량", "선생산반영후수량",
        "실적차감수량", "최종수량", "품번전체실적",
    ]
    if extracted.empty:
        return pd.DataFrame(columns=columns)

    df = extracted.copy()
    df["날짜"] = pd.to_datetime(df["날짜"], errors="coerce").dt.normalize()
    df["수량"] = pd.to_numeric(df["수량"], errors="coerce").fillna(0)
    df["품번"] = df["품번"].apply(normalize_part_no)

    plan = df[df["구분"] == "계획"].copy()
    if start_date is not None:
        plan = plan[plan["날짜"] >= pd.Timestamp(start_date)]
    if end_date is not None:
        plan = plan[plan["날짜"] <= pd.Timestamp(end_date)]
    plan_dates = sorted(plan["날짜"].dropna().unique())
    if not plan_dates:
        return pd.DataFrame(columns=columns)

    baseline = pd.Timestamp(start_date) if start_date is not None else pd.Timestamp(plan_dates[0])
    shortage = df[(df["구분"] == "미달") & (df["날짜"] == baseline)].copy()
    plan_sum = plan.groupby(["품번", "날짜"], sort=False)["수량"].sum().to_dict()
    shortage_sum = shortage.groupby("품번", sort=False)["수량"].sum().to_dict()

    def first_nonempty(series):
        values = series.dropna().astype(str).str.strip()
        values = values[values != ""]
        return values.iloc[0] if not values.empty else ""

    metadata = (
        df.groupby("품번", sort=False)
        .agg({"모델": first_nonempty, "고객사품번": first_nonempty})
        .to_dict("index")
    )
    part_order = plan["품번"].drop_duplicates().tolist()
    records = []
    first_date = pd.Timestamp(plan_dates[0])

    for part_no in part_order:
        plan_values = [float(plan_sum.get((part_no, pd.Timestamp(day)), 0)) for day in plan_dates]
        shortage_value = float(shortage_sum.get(part_no, 0))
        combined = plan_values.copy()
        combined[0] += shortage_value
        after_advance = adjust_row(combined)
        actual_total = float(actuals.get(normalize_part_no(part_no), 0))
        actual_remaining = actual_total
        info = metadata.get(part_no, {})

        for idx, day in enumerate(plan_dates):
            before_actual = float(after_advance[idx])
            actual_used = min(max(before_actual, 0), actual_remaining)
            final_qty = before_actual - actual_used
            actual_remaining -= actual_used
            shortage_for_day = shortage_value if pd.Timestamp(day) == first_date else 0
            if not any((plan_values[idx], shortage_for_day, before_actual, actual_used, final_qty)):
                continue
            records.append({
                "결과시트": output_sheet,
                "품번": part_no,
                "모델": info.get("모델", ""),
                "고객사품번": info.get("고객사품번", ""),
                "날짜": pd.Timestamp(day).strftime("%Y/%m/%d"),
                "계획수량": plan_values[idx],
                "기준일미달": shortage_for_day,
                "미달합산후수량": combined[idx],
                "선생산반영후수량": before_actual,
                "실적차감수량": actual_used,
                "최종수량": final_qty,
                "품번전체실적": actual_total,
            })

    return pd.DataFrame(records, columns=columns)


def _actual_file_path():
    base = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
    return base / "output.xlsx"


def write_result(output_path, sheets):
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        for col in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in col) + 2, 24)
            ws.column_dimensions[col[0].column_letter].width = width

        if (
            ws.max_row >= 2
            and ws.max_column >= 5
            and ws.cell(1, ws.max_column).value == "합계"
        ):
            # MRP 결과는 마지막 행이 합계 행이다. 품번 칸은 비워 두어
            # 후속 처리에서 "합계"가 품번으로 인식되지 않게 한다.
            total_row = ws.max_row
            data_end_row = total_row - 1
            total_col = ws.max_column
            first_date_col = 4
            last_date_col = total_col - 1
            first_date_letter = get_column_letter(first_date_col)
            last_date_letter = get_column_letter(last_date_col)
            total_col_letter = get_column_letter(total_col)

            for row in range(2, data_end_row + 1):
                ws.cell(row, total_col).value = (
                    f"=SUM({first_date_letter}{row}:{last_date_letter}{row})"
                )

            if total_row and data_end_row >= 2:
                for col_idx in range(first_date_col, last_date_col + 1):
                    col_letter = get_column_letter(col_idx)
                    ws.cell(total_row, col_idx).value = (
                        f"=SUM({col_letter}2:{col_letter}{data_end_row})"
                    )
                ws.cell(total_row, total_col).value = (
                    f"=SUM({total_col_letter}2:{total_col_letter}{data_end_row})"
                )
    wb.save(output_path)


def main():
    root = Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        plan_file = select_excel_file("생산계획 파일 선택 (완성/TANK/CORE)", parent=root)
        if not plan_file:
            return

        finish_start_date, finish_end_date = select_date_range(
            root, title="완성공정 추출 기간 선택"
        )
        tank_start_date, tank_end_date = select_date_range(
            root, title="TANK공정 추출 기간 선택"
        )
        core_start_date, core_end_date = select_date_range(
            root, title="CORE공정 추출 기간 선택"
        )
        output_path = select_excel_save_file("MRP 결과 저장", plan_file, "MRP", parent=root)
        if not output_path:
            return

        actual_file = _actual_file_path()
        if not actual_file.exists():
            raise RuntimeError(f"오늘실적 정규화본을 찾지 못했습니다: {actual_file}")
        actuals_by_process = load_actuals_by_process(actual_file)

        output_sheets = {}
        trace_frames = []
        for source_sheet, process, output_sheet in PLAN_SHEETS:
            actuals = actuals_by_process.get(process, actuals_by_process.get("전체", {}))
            extract_options = {"include_accessory": False} if process == "완성" else {}
            if process == "완성":
                start_date, end_date = finish_start_date, finish_end_date
            elif process == "TANK":
                start_date, end_date = tank_start_date, tank_end_date
            else:
                start_date, end_date = core_start_date, core_end_date
            extracted = extract_mrp_source(
                plan_file, source_sheet, process, **extract_options
            )
            output_sheets[output_sheet] = build_mrp_sheet(
                extracted, start_date, end_date, actuals
            )
            trace_frames.append(
                build_calculation_trace(
                    extracted, start_date, end_date, actuals, output_sheet
                )
            )

        output_sheets["계산추적"] = pd.concat(trace_frames, ignore_index=True)

        write_result(output_path, output_sheets)
        messagebox.showinfo("완료", f"MRP 파일을 저장했습니다.\n{output_path}", parent=root)
    except Exception as exc:
        messagebox.showerror("오류", str(exc), parent=root)
    finally:
        root.destroy()


if __name__ == "__main__":
    main()
