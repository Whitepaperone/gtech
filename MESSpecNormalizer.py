from __future__ import annotations

import argparse
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import messagebox

from openpyxl import load_workbook
from openpyxl.styles import Font

from ChangeListParser import cell_text
from CommonUtils import create_progress_window, select_excel_file, select_excel_save_file
from ERPSpecNormalizer import CHANGED_FILL, HEADER_FILL, MANUAL_FILL, normalize_spec


MES_DATA_START_ROW = 2
MES_PART_COL = 1  # A: 품목코드
MES_NAME_COL = 2  # B: 품명
MES_DRAWING_COL = 3  # C: 도번
MES_SPEC_COL = 4  # D: 규격
ORIGINAL_SPEC_HEADER = "원본 규격(확인용)"


def default_output_file(mes_file: str | Path) -> str:
    source = Path(mes_file)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(source.with_name(f"{source.stem}_MES규격정리_{timestamp}.xlsx"))


def unique_sheet_name(workbook, base_name: str) -> str:
    if base_name not in workbook.sheetnames:
        return base_name
    for index in range(2, 1000):
        name = f"{base_name}_{index}"
        if name not in workbook.sheetnames:
            return name
    raise RuntimeError("결과 시트 이름을 만들 수 없습니다.")


def has_cell_fill(cell) -> bool:
    fill = cell.fill
    return bool(fill and fill.fill_type)


def prepare_original_spec_column(ws) -> int:
    original_col = MES_SPEC_COL + 1
    if cell_text(ws.cell(1, original_col).value) != ORIGINAL_SPEC_HEADER:
        ws.insert_cols(original_col)

    header_cell = ws.cell(1, original_col)
    header_cell.value = ORIGINAL_SPEC_HEADER
    header_cell.fill = HEADER_FILL
    header_cell.font = Font(bold=True)
    ws.column_dimensions[header_cell.column_letter].width = 24
    return original_col


def is_mes_spec_sheet(ws) -> bool:
    return cell_text(ws.cell(1, MES_SPEC_COL).value).replace(" ", "") == "규격"


def add_report_sheet(workbook, rows: list[dict[str, str]]) -> None:
    ws = workbook.create_sheet(unique_sheet_name(workbook, "MES규격정리결과"))
    headers = ["상태", "시트", "행", "품목코드", "품명", "도번", "기존 규격", "변경 규격", "사유"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row["status"],
                row["sheet"],
                row["row"],
                row["part_number"],
                row["name"],
                row["drawing"],
                row["old_spec"],
                row["new_spec"],
                row["reason"],
            ]
        )

    for column_cells in ws.columns:
        max_len = max(len(cell_text(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 70)


def normalize_mes_specs(
    mes_file: str | Path,
    output_file: str | Path,
    progress=None,
) -> dict[str, int]:
    workbook = load_workbook(mes_file, keep_vba=Path(mes_file).suffix.lower() == ".xlsm")

    report_rows: list[dict[str, str]] = []
    scanned = changed = manual = ok = skipped_colored = skipped_blank = skipped_sheet = 0

    try:
        target_sheets = [ws for ws in workbook.worksheets if is_mes_spec_sheet(ws)]
        skipped_sheet = len(workbook.worksheets) - len(target_sheets)
        sheet_count = max(len(target_sheets), 1)

        for sheet_index, ws in enumerate(target_sheets, start=1):
            original_spec_col = prepare_original_spec_column(ws)
            max_row = ws.max_row or 1
            for row_index in range(MES_DATA_START_ROW, max_row + 1):
                if progress and row_index % 1000 == 0:
                    sheet_base = 10 + int((sheet_index - 1) / sheet_count * 80)
                    sheet_span = 80 / sheet_count
                    percent = min(90, sheet_base + int((row_index / max(max_row, 1)) * sheet_span))
                    progress(percent, f"MES 규격 정리 중... {ws.title} {row_index:,}/{max_row:,}")

                spec_cell = ws.cell(row_index, MES_SPEC_COL)
                if has_cell_fill(spec_cell):
                    skipped_colored += 1
                    continue

                old_spec = cell_text(spec_cell.value)
                if not old_spec:
                    skipped_blank += 1
                    continue

                scanned += 1
                result = normalize_spec(old_spec)
                if result.status == "skip":
                    skipped_blank += 1
                    continue

                if result.status == "changed":
                    original_cell = ws.cell(row_index, original_spec_col)
                    spec_cell.value = result.value
                    spec_cell.fill = CHANGED_FILL
                    original_cell.value = old_spec
                    original_cell.fill = CHANGED_FILL
                    changed += 1
                    report_rows.append(
                        {
                            "status": "자동변경",
                            "sheet": ws.title,
                            "row": str(row_index),
                            "part_number": cell_text(ws.cell(row_index, MES_PART_COL).value),
                            "name": cell_text(ws.cell(row_index, MES_NAME_COL).value),
                            "drawing": cell_text(ws.cell(row_index, MES_DRAWING_COL).value),
                            "old_spec": old_spec,
                            "new_spec": result.value,
                            "reason": result.reason,
                        }
                    )
                elif result.status == "manual":
                    spec_cell.fill = MANUAL_FILL
                    manual += 1
                    report_rows.append(
                        {
                            "status": "수동확인",
                            "sheet": ws.title,
                            "row": str(row_index),
                            "part_number": cell_text(ws.cell(row_index, MES_PART_COL).value),
                            "name": cell_text(ws.cell(row_index, MES_NAME_COL).value),
                            "drawing": cell_text(ws.cell(row_index, MES_DRAWING_COL).value),
                            "old_spec": old_spec,
                            "new_spec": "",
                            "reason": result.reason,
                        }
                    )
                else:
                    ok += 1

        if progress:
            progress(92, "결과 시트 생성 중...")
        add_report_sheet(workbook, report_rows)
        workbook.save(output_file)
    finally:
        workbook.close()

    if progress:
        progress(100, "완료")

    return {
        "scanned": scanned,
        "changed": changed,
        "manual": manual,
        "ok": ok,
        "skipped_colored": skipped_colored,
        "skipped_blank": skipped_blank,
        "skipped_sheet": skipped_sheet,
        "reported": len(report_rows),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="MES 파일 D열 규격을 ERP 규격 표준화 규칙으로 정리합니다.")
    parser.add_argument("mes_file", nargs="?", help="정리할 MES 엑셀 파일 경로")
    parser.add_argument("-o", "--output", help="저장할 엑셀 경로")
    return parser


def run_cli(mes_file: str, output_file: str | None) -> None:
    output = output_file or default_output_file(mes_file)
    stats = normalize_mes_specs(mes_file, output)
    print(f"저장 완료: {output}")
    print(
        " / ".join(
            [
                f"검사 {stats['scanned']}",
                f"자동변경 {stats['changed']}",
                f"수동확인 {stats['manual']}",
                f"정상 {stats['ok']}",
                f"색칠스킵 {stats['skipped_colored']}",
                f"빈규격스킵 {stats['skipped_blank']}",
            ]
        )
    )


def run_gui() -> None:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        mes_file = select_excel_file("규격을 정리할 MES 엑셀 선택", parent=root)
        if not mes_file:
            return

        output_file = select_excel_save_file("MES 규격 정리 결과 저장", mes_file, "MES규격정리", parent=root)
        if not output_file:
            return

        progress_win, update_progress = create_progress_window(root, "MES 규격 정리")
        try:
            stats = normalize_mes_specs(mes_file, output_file, progress=update_progress)
        finally:
            progress_win.destroy()

        messagebox.showinfo(
            "완료",
            f"MES 규격 정리 파일을 저장했습니다.\n\n"
            f"파일: {output_file}\n"
            f"검사: {stats['scanned']}\n"
            f"자동 변경(노랑): {stats['changed']}\n"
            f"수동 확인(빨강): {stats['manual']}\n"
            f"이미 정상: {stats['ok']}\n"
            f"색칠 스킵: {stats['skipped_colored']}\n"
            f"빈 규격 스킵: {stats['skipped_blank']}\n"
            f"시트 스킵: {stats['skipped_sheet']}",
            parent=root,
        )
    except Exception as exc:
        messagebox.showerror("오류", str(exc), parent=root)
    finally:
        root.destroy()


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.mes_file:
        run_cli(args.mes_file, args.output)
        return

    run_gui()


if __name__ == "__main__":
    main()
