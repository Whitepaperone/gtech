import tempfile
import unittest
from pathlib import Path
from datetime import date
from openpyxl import Workbook, load_workbook
import ShipmentChanges as app

class ShipmentTests(unittest.TestCase):
    def test_group_rows_and_audit_retry(self):
        with tempfile.TemporaryDirectory() as folder:
            source, output, retry = [Path(folder) / name for name in ('source.xlsx', 'output.xlsx', 'retry.xlsx')]
            wb = Workbook()
            ws = wb.active
            ws.title = app.TARGET_SHEET
            ws.append(app.REQUIRED_HEADERS)
            for quantity in (60, 40):
                values = {'DDAY': date(2026, 1, 1), 'SALESCODE': 'S1', 'ITEMCODE': 'P1', 'QTY': quantity}
                ws.append([values.get(app.API_FIELD_BY_HEADER.get(header)) for header in app.REQUIRED_HEADERS])
            wb.save(source)
            records = [dict(DDAY='2026-01-01', SALESCODE='S1', ITEMCODE='P1', QTY=100,
                            CHANGE_TYPE='DELETED', CHANGE_ID='a'*64, OLD_TOTAL_QTY=100, NEW_TOTAL_QTY=0),
                       dict(DDAY='2026-01-02', SALESCODE='S2', ITEMCODE='P2', QTY=10,
                            CHANGE_TYPE='DELETED', CHANGE_ID='b'*64, OLD_TOTAL_QTY=10, NEW_TOTAL_QTY=0)]
            result = app.append_shipment_changes(str(source), str(output), records, mapping_file=None)
            self.assertEqual(result['marked_deleted_count'], 2)
            self.assertEqual(result['not_found_count'], 1)
            wb = load_workbook(output)
            self.assertEqual(wb[app.AUDIT_SHEET].max_row, 3)
            self.assertEqual(wb[app.TARGET_SHEET].cell(2, 7).value, 60)
            self.assertEqual(wb[app.TARGET_SHEET].cell(2, 1).fill.fgColor.rgb[-6:], 'F4CCCC')
            wb.close()
            app.append_shipment_changes(str(output), str(retry), records, mapping_file=None)
            wb = load_workbook(retry)
            self.assertEqual(wb[app.AUDIT_SHEET].max_row, 3)
            wb.close()

    def test_excel_com_group_rows_and_audit_retry(self):
        from unittest.mock import patch
        with patch.object(app, 'append_shipment_changes', app.append_shipment_changes_fast):
            self.test_group_rows_and_audit_retry()

if __name__ == '__main__':
    unittest.main()


