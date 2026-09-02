"""재고실사조정과 생산계획의 품번별 재고를 비교한다."""

from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from AppConstants import STOCK_INVENTORY_COMPARATOR_CONFIG
from CommonUtils import create_progress_window, normalize_part_no
from ExtractPlan import extract_stock_by_part


ADJUSTMENT_COLUMNS = STOCK_INVENTORY_COMPARATOR_CONFIG["adjustment_columns"]
PART_COL_INDEX = column_index_from_string(str(ADJUSTMENT_COLUMNS["part_col"])) - 1
LOT_COL_INDEX = column_index_from_string(str(ADJUSTMENT_COLUMNS["lot_col"])) - 1
LOT_TYPE_HEADER = STOCK_INVENTORY_COMPARATOR_CONFIG["lot_type_header"]
LOT_TYPE_VALUE = STOCK_INVENTORY_COMPARATOR_CONFIG["lot_type_value"]
SYSTEM_QTY_HEADER = STOCK_INVENTORY_COMPARATOR_CONFIG.get("system_qty_header", "전산수량")
DEFECT_LOT_TYPE_VALUE = STOCK_INVENTORY_COMPARATOR_CONFIG.get("defect_lot_type_value", "불량")
WIP_LOT_TYPE_VALUE = STOCK_INVENTORY_COMPARATOR_CONFIG.get("wip_lot_type_value", "재공상태")


def _find_header(columns, expected):
    expected_key = str(expected).replace(" ", "").replace("\n", "")
    for column in columns:
        key = str(column).strip().replace(" ", "").replace("\n", "")
        if key == expected_key:
            return column
    raise RuntimeError(f"재고실사조정 파일에서 '{expected}' 열을 찾지 못했습니다.")


def _lot_date_text(value):
    if pd.isna(value):
        return ""
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value).strftime("%y%m%d")
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(char for char in text if char.isdigit())
    if len(digits) == 8 and digits.startswith("20"):
        return digits[2:]
    return digits.zfill(6) if digits else ""


def read_physical_inventory(adjustment_file, today=None):
    """재고, 불량, 재공 수량을 품번별로 분리해 집계한다."""
    frame = pd.read_excel(adjustment_file, engine="calamine")
    if frame.shape[1] <= max(PART_COL_INDEX, LOT_COL_INDEX):
        required_cols = ", ".join(
            map(str, (ADJUSTMENT_COLUMNS["part_col"], ADJUSTMENT_COLUMNS["lot_col"]))
        )
        raise RuntimeError(f"재고실사조정 파일에 {required_cols}열이 모두 필요합니다.")

    lot_type_col = _find_header(frame.columns, LOT_TYPE_HEADER)
    system_qty_col = _find_header(frame.columns, SYSTEM_QTY_HEADER)
    frame["LOT유형정리"] = frame[lot_type_col].astype(str).str.strip()
    filtered = frame[
        frame["LOT유형정리"].isin(
            {LOT_TYPE_VALUE, DEFECT_LOT_TYPE_VALUE, WIP_LOT_TYPE_VALUE}
        )
    ].copy()

    filtered["품번"] = filtered.iloc[:, PART_COL_INDEX].apply(normalize_part_no)
    filtered["수량"] = pd.to_numeric(
        filtered[system_qty_col].astype(str).str.replace(",", "", regex=False),
        errors="coerce",
    ).fillna(0)
    filtered["로트일자"] = filtered.iloc[:, LOT_COL_INDEX].apply(_lot_date_text)
    filtered = filtered[filtered["품번"] != ""].copy()

    today_text = pd.Timestamp(today or datetime.today()).strftime("%y%m%d")
    actual = (
        filtered[filtered["LOT유형정리"] == LOT_TYPE_VALUE]
        .groupby("품번")["수량"]
        .sum()
        .rename("실사재고수")
    )
    defect_qty = (
        filtered[filtered["LOT유형정리"] == DEFECT_LOT_TYPE_VALUE]
        .groupby("품번")["수량"]
        .sum()
        .rename("불량수량")
    )
    wip_qty = (
        filtered[filtered["LOT유형정리"] == WIP_LOT_TYPE_VALUE]
        .groupby("품번")["수량"]
        .sum()
        .rename("재공수량")
    )
    today_qty = (
        filtered[
            (filtered["LOT유형정리"] == LOT_TYPE_VALUE)
            & (filtered["로트일자"] == today_text)
        ]
        .groupby("품번")["수량"]
        .sum()
        .rename("오늘 생산실적량")
    )
    return pd.concat([actual, defect_qty, wip_qty, today_qty], axis=1).fillna(0).reset_index()


def build_comparison(plan_file, adjustment_file, today=None):
    plan_stock = extract_stock_by_part(plan_file)
    physical_stock = read_physical_inventory(adjustment_file, today=today)
    # 생산계획 품번을 기준으로 병합하여 계획에 없는 실사 품번은 비교에서 제외한다.
    result = plan_stock.merge(physical_stock, on="품번", how="left").fillna(0)

    for column in (
        "실사재고수", "불량수량", "재공수량",
        "생산계획재고수", "오늘 생산실적량",
    ):
        result[column] = pd.to_numeric(result[column], errors="coerce").fillna(0)
    result["오늘생산실적제외수량"] = (
        result["실사재고수"]
        + result["불량수량"]
        + result["재공수량"]
        - result["오늘 생산실적량"]
    )
    result["차이"] = result["오늘생산실적제외수량"] - result["생산계획재고수"]
    result["판정"] = result["차이"].map(
        lambda value: "일치" if value == 0 else ("실사 과다" if value > 0 else "실사 부족")
    )
    result = result.rename(columns={
        "오늘 생산실적량": "오늘생산실적수량",
        "생산계획재고수": "생산계획재고",
    })
    columns = [
        "공정구분", "품번", "실사재고수", "불량수량", "재공수량",
        "오늘생산실적수량", "생산계획재고", "오늘생산실적제외수량", "차이", "판정",
    ]
    result["_공정순서"] = result["공정구분"].map(
        {"완성": 0, "악세서리": 1, "TANK": 2, "코어": 3}
    ).fillna(4)
    return (
        result[columns]
        .assign(_공정순서=result["_공정순서"])
        .sort_values(["_공정순서", "품번"], kind="stable")
        .drop(columns="_공정순서")
        .reset_index(drop=True)
    )


def format_output(output_file):
    workbook = load_workbook(output_file)
    sheet = workbook.active
    sheet.title = "재고비교"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)) + 2
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width, 10), 24)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    workbook.save(output_file)
    workbook.close()


def main():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    plan_file = filedialog.askopenfilename(
        parent=root, title="생산계획 파일 선택", filetypes=[("Excel files", "*.xlsx *.xlsm")]
    )
    if not plan_file:
        return
    adjustment_file = filedialog.askopenfilename(
        parent=root, title="재고실사조정 파일 선택", filetypes=[("Excel files", "*.xlsx *.xls")]
    )
    if not adjustment_file:
        return
    output_file = filedialog.asksaveasfilename(
        parent=root,
        title="재고 비교 결과 저장",
        initialdir=str(Path(adjustment_file).parent),
        initialfile=f"재고실사_생산계획_비교_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx")],
    )
    if not output_file:
        return

    progress_window, progress = create_progress_window(root, "재고 비교 진행")
    try:
        progress(10, "생산계획 재고를 읽는 중...")
        result = build_comparison(plan_file, adjustment_file)
        progress(80, "결과 엑셀을 저장하는 중...")
        result.to_excel(output_file, index=False)
        format_output(output_file)
        progress(100, "완료")
        messagebox.showinfo("완료", f"재고 비교가 완료되었습니다.\n{output_file}", parent=root)
    except Exception as exc:
        messagebox.showerror("오류", str(exc), parent=root)
    finally:
        progress_window.destroy()
        root.destroy()


if __name__ == "__main__":
    main()
