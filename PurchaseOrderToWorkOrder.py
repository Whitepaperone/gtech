# -*- coding: utf-8 -*-
"""구매발주현황상세를 CORE부품 작업지시용 엑셀로 변환한다."""

from __future__ import annotations

import math
import re
import subprocess
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REQUIRED_HEADERS = ("품번", "거래처명", "잔량", "납품요청일")
WORKPLACE = "CORE부품"
TEAM_RULES = (
    ("TUBEASSY", "TUBE ASSY"),
    ("INNERFIN", "INNER FIN"),
    ("TUBEPLATE", "TUBE PLATE절단"),
)


def normalized_text(value: Any) -> str:
    """비교할 때 공백, 밑줄, 하이픈 및 대소문자 차이를 무시한다."""
    return re.sub(r"[\s_\-]+", "", str(value or "")).upper()


def normalize_header(value: Any) -> str:
    """헤더의 공백, 줄바꿈 및 흔한 구분 기호 차이를 제거한다."""
    return re.sub(r"[^0-9A-Za-z가-힣]+", "", str(value or "")).strip()


def classify_supplier(value: Any) -> str | None:
    supplier = normalized_text(value)
    if "지텍" not in supplier:
        return None
    for keyword, team in TEAM_RULES:
        if keyword in supplier:
            return team
    return None


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"\s+.*$", "", text)
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def parse_quantity(value: Any) -> float | int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
    if not math.isfinite(number):
        return None
    return int(number) if number.is_integer() else number


def format_part_number(value: Any) -> str:
    part = str(value or "").strip()
    if part.upper().startswith("S-"):
        return part
    return f"S-{part}" if part else ""


def header_key(value: Any) -> str | None:
    name = normalize_header(value)
    if not name:
        return None
    # 정확히 일치하거나 '품번/품명', '거래처명(Supplier)' 같은 복합 헤더를 허용한다.
    for wanted in REQUIRED_HEADERS:
        if name == wanted or wanted in name:
            return wanted
    return None


def find_header_row(ws, search_rows: int = 200) -> tuple[int, dict[str, int]]:
    wanted = set(REQUIRED_HEADERS)
    last_row = min(ws.max_row, search_rows)
    for row_no in range(1, last_row + 1):
        found: dict[str, int] = {}
        found_rows: dict[str, int] = {}
        # 병합되거나 2~3단 헤더인 경우를 위해 현재 행부터 아래 2개 행을 함께 본다.
        for scan_row in range(row_no, min(row_no + 2, last_row) + 1):
            for cell in ws[scan_row]:
                key = header_key(cell.value)
                if key and key not in found:
                    found[key] = cell.column
                    found_rows[key] = scan_row
        if wanted.issubset(found):
            return max(found_rows.values()), found
    raise ValueError(
        f"시트 '{ws.title}'의 첫 200행에서 필수 헤더를 찾지 못했습니다: "
        + ", ".join(REQUIRED_HEADERS)
    )


def read_orders(input_path: str | Path):
    # 일부 ERP 내보내기 파일은 워크시트 사용 범위를 C열까지만 잘못 기록한다.
    # read_only=True에서는 그 범위를 신뢰해 K/O/W열을 읽지 못하므로 일반 모드로 연다.
    workbook = load_workbook(input_path, data_only=True, read_only=False)
    errors = defaultdict(int)
    totals: defaultdict[tuple[str, str, date], float] = defaultdict(float)
    matched_rows = 0
    source_rows = 0
    sheet_name = ""

    try:
        target = None
        for ws in workbook.worksheets:
            try:
                header_row, columns = find_header_row(ws)
                target = (ws, header_row, columns)
                break
            except ValueError:
                continue
        if target is None:
            raise ValueError(
                "모든 시트를 확인했지만 필수 헤더를 찾지 못했습니다: "
                + ", ".join(REQUIRED_HEADERS)
            )

        ws, header_row, columns = target
        sheet_name = ws.title
        for row_no in range(header_row + 1, ws.max_row + 1):
            source_rows += 1
            supplier = ws.cell(row_no, columns["거래처명"]).value
            team = classify_supplier(supplier)
            if team is None:
                continue

            part_value = ws.cell(row_no, columns["품번"]).value
            part = format_part_number(part_value)
            due = parse_date(ws.cell(row_no, columns["납품요청일"]).value)
            quantity = parse_quantity(ws.cell(row_no, columns["잔량"]).value)
            if not part:
                errors["품번 없음"] += 1
                continue
            if due is None:
                errors["납품요청일 오류"] += 1
                continue
            if quantity is None:
                errors["잔량 오류"] += 1
                continue

            totals[(part, team, due)] += quantity
            matched_rows += 1
    finally:
        workbook.close()

    if not totals:
        raise ValueError(
            "변환할 행이 없습니다. 거래처명이 지텍_Tube Assy, 지텍_Inner Fin, "
            "지텍_Tube Plate 중 하나인지 확인해 주세요."
        )
    return totals, {
        "sheet": sheet_name,
        "source_rows": source_rows,
        "matched_rows": matched_rows,
        "skipped": dict(errors),
    }


def write_work_order(output_path: str | Path, totals) -> tuple[int, int]:
    dates = sorted({due for _, _, due in totals})
    items = sorted({(part, team) for part, team, _ in totals}, key=lambda x: (x[1], x[0]))

    wb = Workbook()
    ws = wb.active
    ws.title = "작업지시"
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(3 + len(dates))}{1 + len(items)}"

    headers: list[Any] = ["품번", "작업장명", "작업반명", *dates]
    ws.append(headers)
    for part, team in items:
        ws.append([part, WORKPLACE, team, *[totals.get((part, team, due), None) for due in dates]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for cell in ws[1][3:]:
        cell.number_format = "yyyy-mm-dd"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    for col in range(4, 4 + len(dates)):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.row_dimensions[1].height = 24

    wb.save(output_path)
    return len(items), len(dates)


def convert(input_path: str | Path, output_path: str | Path) -> dict[str, Any]:
    totals, summary = read_orders(input_path)
    item_count, date_count = write_work_order(output_path, totals)
    summary.update(item_count=item_count, date_count=date_count, output=str(output_path))
    return summary


def run_powershell(script: str) -> str:
    """Windows 기본 대화상자를 표시하고 표준 출력 문자열을 반환한다."""
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    completed = subprocess.run(
        ["powershell.exe", "-NoProfile", "-STA", "-Command", script],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        creationflags=creation_flags,
        check=False,
    )
    return completed.stdout.strip()


def ps_quote(text: str) -> str:
    return "'" + text.replace("'", "''") + "'"


def select_input_file() -> str:
    script = r"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
Add-Type -AssemblyName System.Windows.Forms
$d = New-Object System.Windows.Forms.OpenFileDialog
$d.Title = '구매발주현황상세 엑셀 선택'
$d.Filter = 'Excel 통합 문서 (*.xlsx;*.xlsm)|*.xlsx;*.xlsm'
$d.Multiselect = $false
if ($d.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK) { $d.FileName }
"""
    return run_powershell(script)


def select_output_file(default_name: str) -> str:
    script = f"""
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
Add-Type -AssemblyName System.Windows.Forms
$d = New-Object System.Windows.Forms.SaveFileDialog
$d.Title = '작업지시 엑셀 저장'
$d.Filter = 'Excel 통합 문서 (*.xlsx)|*.xlsx'
$d.DefaultExt = 'xlsx'
$d.AddExtension = $true
$d.FileName = {ps_quote(default_name)}
if ($d.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK) {{ $d.FileName }}
"""
    return run_powershell(script)


def show_message(title: str, message: str, error: bool = False):
    icon = "Error" if error else "Information"
    script = f"""
Add-Type -AssemblyName System.Windows.Forms
[System.Windows.Forms.MessageBox]::Show(
    {ps_quote(message)}, {ps_quote(title)}, 'OK', '{icon}'
) | Out-Null
"""
    run_powershell(script)


def run_gui():
    source = select_input_file()
    if not source:
        return
    output = select_output_file(f"작업지시_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    if not output:
        return
    try:
        result = convert(source, output)
        skipped = sum(result["skipped"].values())
        detail = f"\n제외(값 오류): {skipped}행" if skipped else ""
        show_message(
            "변환 완료",
            f"품번/작업반 {result['item_count']}개, 날짜 {result['date_count']}개\n"
            f"반영된 원본 행: {result['matched_rows']}행{detail}\n\n저장 위치:\n{output}",
        )
    except PermissionError:
        show_message("저장 실패", "출력 파일이 Excel에서 열려 있습니다. 닫고 다시 시도해 주세요.", True)
    except Exception as exc:
        show_message("변환 실패", str(exc), True)


def main():
    if len(sys.argv) == 3:
        result = convert(sys.argv[1], sys.argv[2])
        print(f"완료: {result}")
        return
    run_gui()


if __name__ == "__main__":
    main()
