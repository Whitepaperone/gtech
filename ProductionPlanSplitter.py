"""생산계획 원본 서식을 유지해 한국(CU), AL, TANK조립 계획을 분리 저장한다."""

from copy import copy
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from AppConstants import (
    DATA_START_ROW,
    HEADER_DATE_ROW,
    PRODUCTION_PLAN_SPLITTER_CONFIG,
)
from CommonUtils import (
    build_merged_map,
    create_progress_window,
    normalize_part_no,
    normalize_process_token,
    select_date_range,
)
from ExtractPlan import build_header_lookup, build_mes_date_columns, find_date_start_col


CORE_SHEET = "CORE공정(실적)"
TANK_SHEET = "TANK공정(실적)"


def _config_int(name, default):
    try:
        return int(PRODUCTION_PLAN_SPLITTER_CONFIG.get(name, default))
    except (TypeError, ValueError):
        return default


CORE_FUTURE_PLAN_TOTAL_COL = column_index_from_string(
    str(PRODUCTION_PLAN_SPLITTER_CONFIG.get(
        "core_future_plan_total_col",
        PRODUCTION_PLAN_SPLITTER_CONFIG.get("future_plan_total_col", "AH"),
    ))
)
TANK_FUTURE_PLAN_TOTAL_COL = column_index_from_string(
    str(PRODUCTION_PLAN_SPLITTER_CONFIG.get("tank_future_plan_total_col", "AD"))
)
CORE_EXCLUDE_MARKER_ROW = _config_int("core_exclude_marker_row", 19)
TANK_EXCLUDE_MARKER_ROW = _config_int("tank_exclude_marker_row", 21)
TANK_ASSEMBLY_EXTRA_MARKER_ROW = _config_int(
    "tank_assembly_extra_exclude_marker_row", 20
)
try:
    TANK_ASSEMBLY_EXTRA_FILL_TINT = float(
        PRODUCTION_PLAN_SPLITTER_CONFIG.get(
            "tank_assembly_extra_excluded_fill_tint", -0.15
        )
    )
except (TypeError, ValueError):
    TANK_ASSEMBLY_EXTRA_FILL_TINT = -0.15
OUTPUT_EXCLUDED_COLUMNS = {
    column_index_from_string(str(column).strip())
    for column in PRODUCTION_PLAN_SPLITTER_CONFIG.get(
        "output_excluded_columns", ["A"]
    )
}
EXCLUDED_FILL_THEME = _config_int("excluded_fill_theme", 0)
_configured_tints = PRODUCTION_PLAN_SPLITTER_CONFIG.get("excluded_fill_tints")
if _configured_tints is None:
    _configured_tints = [
        PRODUCTION_PLAN_SPLITTER_CONFIG.get("excluded_fill_tint", -0.35),
    ]
try:
    EXCLUDED_FILL_TINTS = tuple(float(tint) for tint in _configured_tints)
except (TypeError, ValueError):
    EXCLUDED_FILL_TINTS = (-0.35,)
try:
    EXCLUDED_FILL_TINT_TOLERANCE = float(
        PRODUCTION_PLAN_SPLITTER_CONFIG.get("excluded_fill_tint_tolerance", 0.02)
    )
except (TypeError, ValueError):
    EXCLUDED_FILL_TINT_TOLERANCE = 0.02
ROW_EXCLUDE_MARKER_COL = column_index_from_string(
    str(PRODUCTION_PLAN_SPLITTER_CONFIG.get("row_exclude_marker_col", "A"))
)
ROW_EXCLUDED_FILL_THEME = _config_int("row_excluded_fill_theme", 0)
try:
    ROW_EXCLUDED_FILL_TINT = float(
        PRODUCTION_PLAN_SPLITTER_CONFIG.get("row_excluded_fill_tint", -0.50)
    )
except (TypeError, ValueError):
    ROW_EXCLUDED_FILL_TINT = -0.50


def _normalized_process(value) -> str:
    return normalize_process_token(value).upper().replace(" ", "")


def _matches(group: str, value) -> bool:
    process = _normalized_process(value)
    if group == "CU":
        return process == "CU"
    if group == "AL":
        return process.startswith("AL")
    if group == "TANK조립":
        return "TANK조립" in process
    if group == "한국":
        return process == "한국"
    if group == "로봇":
        return "로봇" in process
    if group == "수동":
        return "수동" in process
    if group == "로봇,수동":
        return "로봇" in process or "수동" in process
    return False


def _is_excluded_fill(cell):
    """설정된 '흰색, 배경 1'의 어두운 테마 채우기를 판별한다."""
    fill = cell.fill
    if not fill or fill.fill_type is None:
        return False
    color = fill.fgColor
    if color.type != "theme" or color.theme != EXCLUDED_FILL_THEME:
        return False
    tint = float(color.tint or 0)
    return any(
        abs(tint - excluded_tint) <= EXCLUDED_FILL_TINT_TOLERANCE
        for excluded_tint in EXCLUDED_FILL_TINTS
    )


def _has_theme_fill(cell, theme, tint):
    fill = cell.fill
    if not fill or fill.fill_type is None:
        return False
    color = fill.fgColor
    return (
        color.type == "theme"
        and color.theme == theme
        and abs(float(color.tint or 0) - tint) <= EXCLUDED_FILL_TINT_TOLERANCE
    )


def _is_excluded_row(ws, row):
    """A열이 '흰색, 배경 1, 50% 더 어둡게'인 품번 행인지 판별한다."""
    fill = ws.cell(row, ROW_EXCLUDE_MARKER_COL).fill
    if not fill or fill.fill_type is None:
        return False
    color = fill.fgColor
    if color.type != "theme" or color.theme != ROW_EXCLUDED_FILL_THEME:
        return False
    tint = float(color.tint or 0)
    return abs(tint - ROW_EXCLUDED_FILL_TINT) <= EXCLUDED_FILL_TINT_TOLERANCE


def _excluded_columns(ws, group, date_start_col):
    marker_row = CORE_EXCLUDE_MARKER_ROW if group in {"CU", "AL"} else TANK_EXCLUDE_MARKER_ROW
    excluded = set(OUTPUT_EXCLUDED_COLUMNS)
    # 회색 표시는 고정 정보 열의 불필요 여부만 나타낸다. 날짜 영역까지
    # 검사하면 종료일 계획 열이 빠져 하루 전까지만 보일 수 있다.
    for col in range(1, date_start_col):
        if _is_excluded_fill(ws.cell(marker_row, col)):
            excluded.add(col)
        if (
            group == "TANK조립"
            and _has_theme_fill(
                ws.cell(TANK_ASSEMBLY_EXTRA_MARKER_ROW, col),
                EXCLUDED_FILL_THEME,
                TANK_ASSEMBLY_EXTRA_FILL_TINT,
            )
        ):
            excluded.add(col)
    return excluded


def _selected_columns(ws, group, start_date, end_date):
    """고정 열 + 시작일의 미달(없으면 계획)부터 종료일 계획까지 연속 반환."""
    date_start = find_date_start_col(ws)
    if date_start is None:
        raise RuntimeError(f"{ws.title}: 날짜 시작 열을 찾지 못했습니다.")

    start = pd.Timestamp(start_date).normalize()
    end = pd.Timestamp(end_date).normalize()
    date_columns = build_mes_date_columns(ws)
    start_candidates = [
        item["col"] for item in date_columns
        if item["date"] == start
        and normalize_process_token(item["kind"]).replace(" ", "") == "미달"
    ]
    end_candidates = [
        item["col"] for item in date_columns
        if item["date"] == end
        and normalize_process_token(item["kind"]).replace(" ", "") == "계획"
    ]
    if not start_candidates:
        # 미래 날짜에는 미달 열이 생성되지 않으므로 계획 열부터 시작한다.
        start_candidates = [
            item["col"] for item in date_columns
            if item["date"] == start
            and normalize_process_token(item["kind"]).replace(" ", "") == "계획"
        ]
    if not start_candidates:
        raise RuntimeError(
            f"{ws.title}: {start:%Y-%m-%d}의 미달 또는 계획 열을 찾지 못했습니다."
        )
    if not end_candidates:
        raise RuntimeError(f"{ws.title}: {end:%Y-%m-%d}의 계획 열을 찾지 못했습니다.")

    first_date_col = min(start_candidates)
    last_date_col = max(end_candidates)
    if first_date_col > last_date_col:
        raise RuntimeError(f"{ws.title}: 추출할 날짜 열 범위가 올바르지 않습니다.")

    # 중간 실적 열을 빼지 않고 원본 순서 그대로 연속 복사한다.
    columns = list(range(1, date_start)) + list(range(first_date_col, last_date_col + 1))
    excluded = _excluded_columns(ws, group, date_start)
    deadline_total_col = (
        CORE_FUTURE_PLAN_TOTAL_COL if group in {"CU", "AL"}
        else TANK_FUTURE_PLAN_TOTAL_COL
    )
    excluded.discard(deadline_total_col)
    selected = [col for col in columns if col not in excluded]
    plan_columns = [
        item["col"] for item in date_columns
        if start <= item["date"] <= end
        and normalize_process_token(item["kind"]).replace(" ", "") == "계획"
    ]
    shortage_columns = [
        item["col"] for item in date_columns
        if item["date"] == start
        and normalize_process_token(item["kind"]).replace(" ", "") == "미달"
    ]
    return selected, plan_columns, shortage_columns, date_start


def _has_future_plan_value(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return False
    try:
        return float(value) != 0
    except (TypeError, ValueError):
        return True


def _number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _selected_rows(ws, value_ws, group: str, plan_columns, shortage_columns):
    headers = build_header_lookup(ws)
    process_col = headers.get("process_print_col")
    is_core = group in {"CU", "AL"}
    part_col = headers.get("core_part_col" if is_core else "tank_part_col")
    future_plan_col = (
        CORE_FUTURE_PLAN_TOTAL_COL if is_core else TANK_FUTURE_PLAN_TOTAL_COL
    )
    if not process_col:
        raise RuntimeError(f"{ws.title}: 공정인쇄 열을 찾지 못했습니다.")
    if not part_col:
        raise RuntimeError(f"{ws.title}: 대상 품번 열을 찾지 못했습니다.")

    merged_values = build_merged_map(ws)
    value_merged_values = (
        merged_values if value_ws is ws else build_merged_map(value_ws)
    )
    rows = []
    period_totals = {}
    process_values = {}
    process_matches = 0
    part_matches = 0
    for row in range(DATA_START_ROW, ws.max_row + 1):
        if _is_excluded_row(ws, row):
            continue
        process_value = merged_values.get((row, process_col), ws.cell(row, process_col).value)
        if not _matches(group, process_value):
            continue
        process_matches += 1
        part_value = merged_values.get((row, part_col), ws.cell(row, part_col).value)
        if not normalize_part_no(part_value):
            continue
        part_matches += 1
        plan_total = sum(_number(value_ws.cell(row, col).value) for col in plan_columns)
        shortage_total = sum(
            _number(value_ws.cell(row, col).value) for col in shortage_columns
        )
        # 계획이 없어도 미달 값이 있으면 유지한다. 서로 상쇄되어 합계가
        # 0이어도 원본 값이 존재하는 행은 제거하지 않는다.
        if plan_total == 0 and shortage_total == 0:
            continue
        period_total = plan_total + shortage_total
        rows.append(row)
        period_totals[row] = period_total
        process_values[row] = _normalized_process(process_value)
    # 원본의 날짜/구분 헤더(기본 33~34행)부터 시작하고 상단 보고서 영역은 제외한다.
    return (
        list(range(HEADER_DATE_ROW, DATA_START_ROW)) + rows,
        period_totals,
        process_values,
        future_plan_col,
        part_col,
    )


def _copy_sheet(
    source, value_source, target, row_numbers, column_numbers, value_overrides=None
):
    value_overrides = value_overrides or {}
    row_map = {old: new for new, old in enumerate(row_numbers, 1)}
    col_map = {old: new for new, old in enumerate(column_numbers, 1)}

    for old_row, new_row in row_map.items():
        source_dim = source.row_dimensions[old_row]
        target_dim = target.row_dimensions[new_row]
        target_dim.height = source_dim.height
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel

        for old_col, new_col in col_map.items():
            source_cell = source.cell(old_row, old_col)
            target_cell = target.cell(new_row, new_col)
            # 수식 자체가 아니라 원본 파일에 저장된 계산 결과값만 복사한다.
            target_cell.value = value_overrides.get(
                (old_row, old_col), value_source.cell(old_row, old_col).value
            )
            if source_cell.has_style:
                # 다른 통합문서의 내부 style_id를 그대로 복사하면 저장 시
                # "list index out of range"가 발생할 수 있으므로 구성요소별로
                # 할당해 대상 통합문서의 스타일 테이블에 정상 등록시킨다.
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.number_format = source_cell.number_format
                target_cell.protection = copy(source_cell.protection)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for old_col, new_col in col_map.items():
        old_letter = get_column_letter(old_col)
        new_letter = get_column_letter(new_col)
        source_dim = source.column_dimensions[old_letter]
        target_dim = target.column_dimensions[new_letter]
        target_dim.width = source_dim.width
        target_dim.hidden = source_dim.hidden
        target_dim.outlineLevel = source_dim.outlineLevel

    for merged in source.merged_cells.ranges:
        rows = range(merged.min_row, merged.max_row + 1)
        cols = range(merged.min_col, merged.max_col + 1)
        if all(r in row_map for r in rows) and all(c in col_map for c in cols):
            target.merge_cells(
                start_row=row_map[merged.min_row], start_column=col_map[merged.min_col],
                end_row=row_map[merged.max_row], end_column=col_map[merged.max_col],
            )

    if source.freeze_panes:
        frozen = source[source.freeze_panes] if isinstance(source.freeze_panes, str) else source.freeze_panes
        mapped_row = row_map.get(frozen.row)
        mapped_col = col_map.get(frozen.column)
        if mapped_row and mapped_col:
            target.freeze_panes = target.cell(mapped_row, mapped_col)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.sheet_format = copy(source.sheet_format)
    target.sheet_state = "visible"
    header_row = row_map[HEADER_DATE_ROW]
    if target.max_row > header_row and target.max_column:
        target.auto_filter.ref = f"A{header_row}:{get_column_letter(target.max_column)}{target.max_row}"
    return row_map, col_map


def _append_totals(
    target, data_rows, sum_source_cols, col_map, part_col, process_values, group
):
    groups = []
    if group == "AL":
        groups = [
            ("AL 합계", [row for row in data_rows if process_values.get(row) == "AL"]),
            ("AL 대 합계", [row for row in data_rows if process_values.get(row) == "AL대"]),
        ]
        groups = [(label, rows) for label, rows in groups if rows]
        if data_rows:
            groups.append(("AL 총합계", data_rows))
    elif group == "로봇,수동":
        groups = [
            ("로봇 합계", [row for row in data_rows if "로봇" in process_values.get(row, "")]),
            ("수동 합계", [row for row in data_rows if "수동" in process_values.get(row, "")]),
        ]
        groups = [(label, rows) for label, rows in groups if rows]
    else:
        groups = [("합계", data_rows)]

    label_col = col_map.get(part_col, 1)
    target_row_by_source = {row: 3 + index for index, row in enumerate(data_rows)}
    for label, rows in groups:
        target_row = target.max_row + 1
        target.cell(target_row, label_col, label)
        for old_col in sum_source_cols:
            new_col = col_map.get(old_col)
            if not new_col:
                continue
            total = sum(
                _number(target.cell(target_row_by_source[row], new_col).value)
                for row in rows
            )
            target.cell(target_row, new_col, total)
        for col in range(1, target.max_column + 1):
            cell = target.cell(target_row, col)
            cell.font = Font(name="맑은 고딕", size=10, bold=True)
            cell.fill = PatternFill("solid", fgColor="FFD9E1F2")
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _format_for_print(
    target, date_start_col, col_map, model_col, finish_part_col, group
):
    for old_col, new_col in col_map.items():
        letter = get_column_letter(new_col)
        values = [
            str(target.cell(row, new_col).value or "").replace("\n", " ")
            for row in range(1, target.max_row + 1)
        ]
        content_width = max((len(value) for value in values), default=4) + 1
        if old_col >= date_start_col:
            # 날짜/수량 열은 화면에서도 촘촘하게 보이도록 고정 폭에 가깝게 제한한다.
            width = min(max(content_width, 5), 7)
        else:
            width = min(max(content_width, 5), 15)
        target.column_dimensions[letter].width = width

    for row in range(1, target.max_row + 1):
        for col in range(1, target.max_column + 1):
            cell = target.cell(row, col)
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
                shrink_to_fit=True,
            )

            if row >= 3 and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                font = copy(cell.font)
                font.sz = max(float(font.sz or 10), 12)
                cell.font = font

    # 모델명 열은 실제 내용에 맞춰 AutoFit에 가까운 너비를 적용한다.
    model_target_col = col_map.get(model_col)
    if model_target_col:
        model_values = [
            str(target.cell(row, model_target_col).value or "").replace("\n", " ")
            for row in range(1, target.max_row + 1)
        ]
        model_width = max((len(value) for value in model_values), default=8) + 2
        target.column_dimensions[get_column_letter(model_target_col)].width = min(
            max(model_width, 10), 35
        )

    # 완성품번 열도 실제 품번 길이에 맞춰 AutoFit한다.
    finish_part_target_col = col_map.get(finish_part_col)
    if finish_part_target_col:
        part_values = [
            str(target.cell(row, finish_part_target_col).value or "").replace("\n", " ")
            for row in range(1, target.max_row + 1)
        ]
        part_width = max((len(value) for value in part_values), default=8) + 2
        target.column_dimensions[get_column_letter(finish_part_target_col)].width = min(
            max(part_width, 10), 35
        )

    # 사용자가 지정한 숫자 영역은 숫자 길이에 맞춰 다시 넓힌다.
    numeric_end_col = 11 if group == "TANK조립" else 17  # H~K / H~Q
    for col in range(8, min(numeric_end_col, target.max_column) + 1):
        numeric_values = [
            str(target.cell(row, col).value)
            for row in range(3, target.max_row + 1)
            if isinstance(target.cell(row, col).value, (int, float))
            and not isinstance(target.cell(row, col).value, bool)
        ]
        if numeric_values:
            numeric_width = max(len(value) for value in numeric_values) + 2
            target.column_dimensions[get_column_letter(col)].width = min(
                max(numeric_width, 7), 13
            )

    # 백분율은 내부 소수 자릿수가 길어도 표시값은 짧으므로 고정 폭을 쓴다.
    for col in range(1, target.max_column + 1):
        header_texts = " ".join(
            str(target.cell(row, col).value or "") for row in (1, 2)
        ).replace(" ", "")
        has_percent_format = any(
            "%" in str(target.cell(row, col).number_format or "")
            for row in range(3, target.max_row + 1)
        )
        if "실적율" in header_texts or "실적률" in header_texts or has_percent_format:
            target.column_dimensions[get_column_letter(col)].width = 9
    target.sheet_properties.pageSetUpPr.fitToPage = True
    target.page_setup.orientation = "landscape"
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 0
    target.print_options.horizontalCentered = True
    target.sheet_view.zoomScale = 75


def _add_extract(
    source_wb, value_wb, output_wb, source_sheet, group, title, start_date, end_date
):
    if source_sheet not in source_wb.sheetnames:
        raise RuntimeError(f"필수 시트가 없습니다: {source_sheet}")
    source = source_wb[source_sheet]
    value_source = value_wb[source_sheet]
    header_columns = build_header_lookup(source)
    model_col = header_columns.get("model_col")
    finish_part_col = header_columns.get("finish_part_col")
    columns, plan_columns, shortage_columns, date_start_col = _selected_columns(
        source, group, start_date, end_date
    )
    rows, period_totals, process_values, future_plan_col, part_col = _selected_rows(
        source, value_source, group, plan_columns, shortage_columns
    )
    target = output_wb.create_sheet(title=title)
    overrides = {
        (row, future_plan_col): total for row, total in period_totals.items()
    }
    overrides[(HEADER_DATE_ROW, future_plan_col)] = "기한내 수량"
    row_map, col_map = _copy_sheet(
        source, value_source, target, rows, columns, value_overrides=overrides
    )
    data_rows = [row for row in rows if row >= DATA_START_ROW]
    _append_totals(
        target, data_rows,
        [future_plan_col] + [col for col in columns if col >= date_start_col],
        col_map,
        part_col, process_values, group,
    )
    _format_for_print(
        target, date_start_col, col_map, model_col, finish_part_col, group
    )


def _new_workbook(source_workbook):
    wb = Workbook()
    wb.remove(wb.active)
    # 테마 색을 사용하는 셀이 원본과 똑같은 색으로 렌더링되도록 이식한다.
    wb.loaded_theme = source_workbook.loaded_theme
    return wb


def extract_split_plan(
    plan_file,
    start_date,
    end_date,
    tank_assembly_start_date=None,
    tank_assembly_end_date=None,
    output_dir=None,
    progress=None,
):
    source_path = Path(plan_file)
    if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise RuntimeError("서식 복사는 .xlsx 또는 .xlsm 파일만 지원합니다.")
    output_dir = Path(output_dir) if output_dir else source_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    prefix = pd.Timestamp(start_date).strftime("%Y%m%d")
    tank_assembly_start_date = tank_assembly_start_date or start_date
    tank_assembly_end_date = tank_assembly_end_date or end_date
    korea_path = output_dir / f"{prefix} 생산계획-한국.xlsx"
    combined_path = output_dir / f"{prefix} 생산계획 AL, TANK조립.xlsx"

    if progress:
        progress(5, "생산계획 파일을 불러오는 중...")
    # data_only 통합문서에도 셀 서식은 그대로 들어 있으므로 파일을 두 번
    # 열 필요가 없다. 저장된 계산값과 서식을 한 번의 로딩으로 함께 사용한다.
    source_wb = load_workbook(source_path, data_only=True)
    value_wb = source_wb
    korea_wb = _new_workbook(source_wb)
    combined_wb = _new_workbook(source_wb)

    if progress:
        progress(25, "CORE공정의 CU 계획을 추출하는 중...")
    _add_extract(source_wb, value_wb, korea_wb, CORE_SHEET, "CU", "CU", start_date, end_date)

    if progress:
        progress(42, "TANK공정의 한국 계획을 추출하는 중...")
    _add_extract(
        source_wb, value_wb, korea_wb,
        TANK_SHEET, "한국", "TANK-한국", start_date, end_date,
    )

    if progress:
        progress(58, "CORE공정의 AL 계획을 추출하는 중...")
    _add_extract(source_wb, value_wb, combined_wb, CORE_SHEET, "AL", "AL", start_date, end_date)

    if progress:
        progress(72, "TANK 로봇, 수동 계획을 추출하는 중...")
    _add_extract(
        source_wb, value_wb, combined_wb,
        TANK_SHEET, "로봇,수동", "로봇,수동", start_date, end_date,
    )

    if progress:
        progress(82, "TANK조립 계획을 추출하는 중...")
    _add_extract(
        source_wb, value_wb, combined_wb,
        TANK_SHEET, "TANK조립", "TANK조립",
        tank_assembly_start_date, tank_assembly_end_date,
    )

    if progress:
        progress(90, "한국(CU) 생산계획 파일을 저장하는 중...")
    korea_wb.save(korea_path)

    if progress:
        progress(96, "AL, TANK조립 생산계획 파일을 저장하는 중...")
    combined_wb.save(combined_path)

    if progress:
        progress(100, "생산계획 추출이 완료되었습니다.")
    return korea_path, combined_path


def main():
    root = tk.Tk()
    root.withdraw()
    plan_file = filedialog.askopenfilename(
        parent=root, title="생산계획 파일 선택",
        filetypes=[("Excel files", "*.xlsx *.xlsm")],
    )
    if not plan_file:
        return
    start_date, end_date = select_date_range(root, title="추출 기간 선택")
    if start_date is None or end_date is None:
        messagebox.showinfo("취소", "추출 기간을 선택하지 않았습니다.", parent=root)
        return
    tank_start_date, tank_end_date = select_date_range(
        root, title="TANK조립 추출 기간 선택"
    )
    if tank_start_date is None or tank_end_date is None:
        messagebox.showinfo(
            "취소", "TANK조립 추출 기간을 선택하지 않았습니다.", parent=root
        )
        return
    progress_win, update_progress = create_progress_window(root, title="생산계획 추출 중")
    try:
        korea_path, combined_path = extract_split_plan(
            plan_file,
            start_date,
            end_date,
            tank_assembly_start_date=tank_start_date,
            tank_assembly_end_date=tank_end_date,
            progress=update_progress,
        )
    except Exception as exc:
        messagebox.showerror("오류", str(exc), parent=root)
        return
    finally:
        try:
            progress_win.destroy()
        except Exception:
            pass
    messagebox.showinfo("완료", f"저장 완료\n{korea_path}\n{combined_path}", parent=root)


if __name__ == "__main__":
    main()
