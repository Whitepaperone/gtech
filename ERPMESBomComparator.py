"""ERP 기준 MES BOM 비교. 실행: python ERPMESBomComparator.py"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import queue
import re
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

# 이 작업 폴더에 설치된 openpyxl도 지원한다.
local_packages = Path(__file__).resolve().parent / '.local-python'
if local_packages.is_dir():
    sys.path.insert(0, str(local_packages))
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

GREEN = 'E2F0D9'
ORANGE = 'FCE4D6'
YELLOW = 'FFF2CC'


def text(value):
    return '' if value is None else str(value).strip()


def part_key(value):
    return text(value).split('-', 1)[0].strip()


def is_outsourcing_cost(value):
    return '-(외주비)' in re.sub(r'\s+', '', text(value))


def quantity(value):
    try:
        result = Decimal(text(value).replace(',', ''))
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def level_zero(value):
    # 0, 0.0, '0', '.0' 등의 표기 지원.
    raw = text(value)
    return bool(raw) and bool(re.fullmatch(r'\.*0(?:\.0+)?', raw))


def header(ws, qty_name):
    required = {'레벨', '품번', qty_name}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 100)):
        names = {}
        for cell in row:
            key = re.sub(r'\s+', '', text(cell.value))
            if key in required:
                if key in names:
                    raise ValueError(f'{ws.title}: {cell.row}행에 {key} 헤더가 중복됩니다.')
                names[key] = cell.column
        if required <= names.keys():
            return row[0].row, names
    raise ValueError(f'{ws.title}: 첫 100행에서 레벨, 품번, {qty_name} 헤더를 찾지 못했습니다.')


def collect(ws, qty_name):
    start, cols = header(ws, qty_name)
    groups = defaultdict(list)
    issues = []
    for row in ws.iter_rows(min_row=start + 1):
        original = row[cols['품번'] - 1].value
        level = row[cols['레벨'] - 1].value
        if level_zero(level) or is_outsourcing_cost(original):
            continue
        if not text(original):
            if text(level):
                issues.append([ws.title, row[0].row, '', '품번 누락'])
            continue
        key = part_key(original)
        if not key:
            raise ValueError(f'{ws.title} {row[0].row}행: 접미사 제거 후 품번이 비었습니다.')
        value = row[cols[qty_name] - 1].value
        qty = quantity(value)
        if qty is None:
            issues.append([ws.title, row[0].row, text(original), f'{qty_name} 확인 필요: {text(value) or "빈칸 또는 계산값 없는 수식"}'])
        groups[key].append((row[0].row, text(original), qty))
    if not groups and not any(is_outsourcing_cost(ws.cell(r, cols['품번']).value)
                              for r in range(start + 1, ws.max_row + 1)):
        raise ValueError(f'{ws.title}: 0레벨 제외 후 비교할 품번이 없습니다.')
    return start, cols, groups, issues


def total(items):
    return None if any(item[2] is None for item in items) else sum((item[2] for item in items), Decimal(0))


def number(value):
    return None if value is None else float(value)


def append_safe(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def style_table(ws, header_row=1):
    ws.freeze_panes = f'A{header_row + 1}'
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}'
    for cell in ws[header_row]:
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.font = Font(color='FFFFFF', bold=True)
    for col in ws.columns:
        width = max((len(text(c.value)) for c in col[:200]), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(48, max(12, width * 1.5))


def compare(erp_path, mes_path, output_path, erp_sheet=None, mes_sheet=None):
    target = Path(output_path).resolve()
    if target in {Path(erp_path).resolve(), Path(mes_path).resolve()}:
        raise ValueError('결과 파일은 원본과 다른 경로로 저장해야 합니다.')
    erp_book = load_workbook(erp_path, data_only=True)
    try:
        mes_book = load_workbook(mes_path, data_only=True)
        try:
            erp = erp_book[erp_sheet] if erp_sheet else erp_book.active
            mes = mes_book[mes_sheet] if mes_sheet else mes_book.active
            erp_header, erp_cols, eg, ei = collect(erp, '수량')
            _, _, mg, mi = collect(mes, '소요량')
            book = Workbook()
            result = book.active
            result.title = 'ERP 기준 비교'
            for row in erp.iter_rows(values_only=True):
                append_safe(result, list(row))
            result.insert_cols(8, 6)
            offset = 7
            part_column = erp_cols['품번'] + (6 if erp_cols['품번'] >= 8 else 0)
            extras = ['비교 품번', 'ERP 합계수량', 'MES 합계소요량', '차이(MES-ERP)', '비교결과', 'MES 원래 품번']
            for col, name in enumerate(extras, offset + 1):
                result.cell(erp_header, col, name)
            counts = defaultdict(int)
            for key, entries in eg.items():
                eq = total(entries)
                mq = total(mg[key]) if key in mg else None
                if key not in mg:
                    status, color = 'MES 미구성', None
                elif eq is None or mq is None:
                    status, color = '수량 확인 필요', YELLOW
                elif abs(eq - mq) <= Decimal('0.000001'):
                    status, color = '수량 일치', GREEN
                else:
                    status, color = '수량 불일치', ORANGE
                counts[status] += 1
                delta = mq - eq if eq is not None and mq is not None else None
                originals = ', '.join(dict.fromkeys(i[1] for i in mg.get(key, [])))
                for row_no, _, _ in entries:
                    for col, value in enumerate([key, number(eq), number(mq), number(delta), status, originals], offset + 1):
                        cell = result.cell(row_no, col, value)
                        if isinstance(value, str):
                            cell.data_type = 's'
                    if color:
                        result.cell(row_no, part_column).fill = PatternFill('solid', fgColor=color)
                        result.cell(row_no, offset + 5).fill = PatternFill('solid', fgColor=color)
            # 비교 기록은 원본 행 번호로 작성한 뒤, 제외 행을 아래부터 삭제한다.
            for row_no in range(erp.max_row, erp_header, -1):
                if is_outsourcing_cost(erp.cell(row_no, erp_cols['품번']).value):
                    result.delete_rows(row_no)
            only = book.create_sheet('MES에만 있는 품번')
            append_safe(only, ['비교 품번', 'MES 원래 품번', 'MES 합계소요량', 'MES 엑셀 행', '비교결과'])
            for key in sorted(mg.keys() - eg.keys()):
                entries = mg[key]
                append_safe(only, [key, ', '.join(dict.fromkeys(i[1] for i in entries)), number(total(entries)), ', '.join(str(i[0]) for i in entries), 'ERP에 없음'])
            counts['MES에만 있음'] = len(mg.keys() - eg.keys())
            issues = book.create_sheet('확인 필요')
            append_safe(issues, ['구분', '원본 시트', '엑셀 행', '원래 품번', '사유'])
            for source, entries in [('ERP', ei), ('MES', mi)]:
                for entry in entries:
                    append_safe(issues, [source, *entry])
            summary = book.create_sheet('비교 안내', 0)
            for row in [
                ['항목', '내용'], ['ERP 파일', str(erp_path)], ['ERP 시트', erp.title],
                ['MES 파일', str(mes_path)], ['MES 시트', mes.title],
                ['비교 기준', '양쪽 0레벨 제외. 나머지 품번은 첫 - 포함 뒷부분 제거 후 비교. 원본 품번 유지.'],
                ['외주비 제외', '-(외주비)가 붙은 품번은 양쪽 비교/수량 합산에서 제외하고 ERP 결과 행에서도 삭제.'],
                ['수량 기준', '같은 비교 품번의 수량/소요량 단순 합산. 상위 수량은 곱하지 않음.'],
                ['색상', '초록: 수량 일치 / 주황: 수량 불일치 / 노랑: 수량 확인 필요 / 무색: MES 미구성'],
                ['중복 품번', '각 ERP 행에 품번별 합계를 반복 표시. 합계수량 열을 다시 합산하지 마세요.'],
                ['주의', '완제품 한 종류의 동일 기준 수량 BOM끼리 비교하세요. 레벨/공정 구조는 비교하지 않습니다.'],
                ['출력 방식', 'ERP 외주비 행 삭제. 비교 결과 6개 열을 H열 앞에 삽입(H~M). 원본 서식과 수식은 복제하지 않습니다.'],
                ['수량 허용오차', '0.000001'], ['집계 단위', '접미사 제거 후 고유 품번 수'],
                *[[k, counts[k]] for k in ['수량 일치', '수량 불일치', '수량 확인 필요', 'MES 미구성', 'MES에만 있음']],
                ['확인 필요 원본 행 수', len(ei) + len(mi)],
            ]:
                append_safe(summary, row)
            for ws in book:
                style_table(ws, erp_header if ws is result else 1)
            summary.column_dimensions['B'].width = 105
            book.save(target)
            return dict(counts)
        finally:
            mes_book.close()
    finally:
        erp_book.close()


def gui():
    root = tk.Tk()
    root.title('ERP · MES BOM 비교')
    root.geometry('780x350')
    root.columnconfigure(1, weight=1)
    paths, sheets, combos = {}, {}, {}

    def choose(kind):
        filename = filedialog.askopenfilename(title=f'{kind} BOM 선택', filetypes=[('Excel', '*.xlsx *.xlsm')])
        if not filename:
            return
        try:
            wb = load_workbook(filename, read_only=True)
            try:
                names, active = wb.sheetnames, wb.active.title
            finally:
                wb.close()
            paths[kind].set(filename)
            combos[kind]['values'] = names
            sheets[kind].set(active)
        except Exception as exc:
            messagebox.showerror('파일 열기 실패', str(exc))

    for index, kind in enumerate(['ERP', 'MES']):
        paths[kind], sheets[kind] = tk.StringVar(), tk.StringVar()
        ttk.Label(root, text=f'{kind} BOM').grid(row=index * 2, column=0, padx=12, pady=12)
        ttk.Entry(root, textvariable=paths[kind], state='readonly').grid(row=index * 2, column=1, sticky='ew')
        ttk.Button(root, text='파일 선택', command=lambda k=kind: choose(k)).grid(row=index * 2, column=2, padx=12)
        combos[kind] = ttk.Combobox(root, textvariable=sheets[kind], state='readonly')
        combos[kind].grid(row=index * 2 + 1, column=1, sticky='ew')
    ttk.Label(root, text='0레벨 제외 · -(외주비) 품번 제외 · 첫 - 이후 제거 · 같은 품번 수량 합산\n초록: 일치 / 주황: 수량 불일치 / 무색: MES 미구성\n수량은 엑셀 값 그대로 합산합니다. 상위 수량을 곱하지 않습니다.\n지원 파일: .xlsx, .xlsm (.xls는 Excel에서 .xlsx로 저장 후 선택)', justify='left').grid(row=4, column=0, columnspan=3, padx=12, pady=18, sticky='w')
    status = tk.StringVar(value='ERP와 MES 파일 및 비교할 시트를 선택하세요.')
    messages = queue.Queue()

    def start():
        if not all(v.get() for v in paths.values()):
            messagebox.showwarning('파일 선택', 'ERP와 MES 파일을 모두 선택하세요.')
            return
        output = filedialog.asksaveasfilename(title='비교 결과 저장', defaultextension='.xlsx', initialfile=f'ERP_MES_BOM_비교_{datetime.now():%Y%m%d_%H%M%S}.xlsx', filetypes=[('Excel', '*.xlsx')])
        if not output:
            return
        args = (paths['ERP'].get(), paths['MES'].get(), output, sheets['ERP'].get(), sheets['MES'].get())
        button.config(state='disabled')
        status.set('비교 중입니다…')

        def work():
            try:
                messages.put((True, compare(*args), output))
            except Exception as exc:
                messages.put((False, str(exc), output))
        threading.Thread(target=work, daemon=True).start()
        root.after(100, poll)

    def poll():
        try:
            success, result, output = messages.get_nowait()
        except queue.Empty:
            root.after(100, poll)
            return
        button.config(state='normal')
        status.set('저장 완료' if success else '비교 실패')
        if success:
            messagebox.showinfo('비교 완료', '\n'.join(f'{k}: {v}품번' for k, v in result.items()) + f'\n\n저장: {output}')
        else:
            messagebox.showerror('비교 실패', result)

    button = ttk.Button(root, text='비교하고 결과 저장', command=start)
    button.grid(row=5, column=1, pady=8)
    ttk.Label(root, textvariable=status).grid(row=6, column=0, columnspan=3, padx=12, sticky='w')
    root.mainloop()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--erp')
    parser.add_argument('--mes')
    parser.add_argument('--output')
    parser.add_argument('--erp-sheet')
    parser.add_argument('--mes-sheet')
    args = parser.parse_args()
    if any([args.erp, args.mes, args.output]):
        if not all([args.erp, args.mes, args.output]):
            parser.error('--erp, --mes, --output을 함께 지정하세요.')
        print(compare(args.erp, args.mes, args.output, args.erp_sheet, args.mes_sheet))
    else:
        gui()
