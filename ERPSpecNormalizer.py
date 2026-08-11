from __future__ import annotations

import argparse
import re
import sqlite3
import tkinter as tk
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from ChangeListParser import cell_text
from CommonUtils import create_progress_window, select_excel_file, select_excel_save_file
from ERPPartsDb import (
    DEFAULT_ERP_DB,
    ERP_DATA_START_ROW,
    ERP_FIXED_COLUMNS,
    create_erp_db,
    normalize_part_lookup,
)


CHECKED_FILL = PatternFill(fill_type="solid", fgColor="E2F0D9")
CHANGED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
MANUAL_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")

NUMBER_PATTERN = r"(?:\d+(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
NUMBER_RE = re.compile(NUMBER_PATTERN)
DIM_SEPARATOR_RE = re.compile(r"[xX*×]")
UNIT_TOKEN_RE = re.compile(
    rf"(?:(?<![A-Za-z0-9.,])|(?<=[xX]))(?:(?P<num1>{NUMBER_PATTERN})\s*(?P<unit1>[THWL])(?=\s*(?:[xX*×]|$|[^A-Za-z0-9.,]))|(?P<unit2>[THWL])\s*(?P<num2>{NUMBER_PATTERN})(?=\s*(?:[xX*×]|$|[^A-Za-z0-9.,])))",
    re.IGNORECASE,
)
THICKNESS_SUFFIX_RE = re.compile(
    rf"(?P<thick>{NUMBER_PATTERN})\s*(?P<label>[HT])(?=\s*(?:[xX*×]|$|[_\-/(),;:]))",
    re.IGNORECASE,
)
DIMENSION_TAIL_RE = re.compile(
    rf"^\s*[xX*×]\s*(?P<width>{NUMBER_PATTERN})(?P<width_unit>\s*W)?"
    rf"(?:\s*[xX*×]\s*(?P<length>{NUMBER_PATTERN})(?P<length_unit>\s*L)?)?"
    r"(?P<after>.*)$",
    re.IGNORECASE,
)
DUAL_THWL_RE = re.compile(
    rf"^\s*(?P<thick>{NUMBER_PATTERN})\s*T\s*[xX*×]\s*"
    rf"(?P<height>{NUMBER_PATTERN})\s*H\s*[xX*×]\s*"
    rf"(?P<width>{NUMBER_PATTERN})\s*W\s*[xX*×]\s*"
    rf"(?P<length>{NUMBER_PATTERN})\s*L(?P<suffix>.*)$",
    re.IGNORECASE,
)
SKIP_CATEGORY_EXACT = {
    "",
    "0",
    "가공품",
    "기능품",
    "기타",
    "부재료",
    "분류",
    "분류없음",
    "사용",
    "상품",
    "소모품",
    "외주비용",
    "제관품",
}
SKIP_CATEGORY_PREFIXES = ("반제품", "완제품")


@dataclass(frozen=True)
class NormalizeResult:
    status: str
    value: str
    reason: str = ""


def load_db_part_keys(db_path: str | Path = DEFAULT_ERP_DB) -> set[str]:
    if not Path(db_path).exists():
        raise FileNotFoundError(f"ERP DB가 없습니다: {db_path}")

    create_erp_db(db_path)
    with sqlite3.connect(db_path) as conn:
        rows = conn.execute("SELECT part_key FROM erp_parts WHERE part_key <> ''").fetchall()

    part_keys = {row[0] for row in rows if row[0]}
    if not part_keys:
        raise ValueError("ERP DB에 저장된 품번이 없습니다. ERPDbUpdater로 먼저 DB를 갱신해 주세요.")
    return part_keys


def latest_db_source_file(db_path: str | Path = DEFAULT_ERP_DB) -> str:
    if not Path(db_path).exists():
        return ""

    create_erp_db(db_path)
    with sqlite3.connect(db_path) as conn:
        row = conn.execute(
            """
            SELECT source_file
            FROM erp_parts
            WHERE COALESCE(source_file, '') <> ''
            GROUP BY source_file
            ORDER BY MAX(updated_at) DESC, COUNT(*) DESC
            LIMIT 1
            """
        ).fetchone()
    return row[0] if row and row[0] else ""


def format_number(text: str) -> str:
    normalized_text = text.replace(",", "")
    sign = ""
    if normalized_text.startswith("-"):
        sign = "-"
        normalized_text = normalized_text[1:]
    integer_text, _, decimal_part = normalized_text.partition(".")
    integer_value = int(integer_text or "0")

    if integer_value == 0:
        return f"{sign}{normalized_text}"
    if integer_value >= 10 and decimal_part and set(decimal_part) == {"0"}:
        return f"{sign}{integer_value}"
    if decimal_part:
        return f"{sign}{integer_value}.{decimal_part}"
    if 1 <= integer_value <= 9:
        return f"{sign}{integer_value}.0"
    return f"{sign}{integer_value}"


def format_length_number(text: str) -> str:
    return format_number(text)


def format_optional_decimal_number(text: str) -> str:
    return format_number(text)


def build_spec_value(label: str, thick: str, width: str, length: str | None = None, desc: str = "") -> str:
    value = f"{format_number(thick)}{label.upper()}x{format_number(width)}W"
    if length:
        value = f"{value}x{format_length_number(length)}L"
    if desc:
        value = f"{value}_{desc}"
    return value


def build_unlabeled_spec_value(label: str, thick: str, numbers: list[str], desc: str = "") -> str:
    parts = [f"{format_number(thick)}{label.upper()}"]
    parts.extend(format_optional_decimal_number(number) for number in numbers)
    value = "x".join(parts)
    if desc:
        value = f"{value}_{desc}"
    return value


def build_thwl_spec_value(match: re.Match[str]) -> str:
    return (
        f"{format_number(match.group('thick'))}T"
        f"x{format_number(match.group('height'))}H"
        f"x{format_number(match.group('width'))}W"
        f"x{format_optional_decimal_number(match.group('length'))}L"
        f"{match.group('suffix') or ''}"
    )


def normalize_dimension_separators(text: str) -> str:
    return text.replace("*", "x")


def unit_token_value(match: re.Match[str]) -> tuple[str, str]:
    unit = match.group("unit1") or match.group("unit2") or ""
    number = match.group("num1") or match.group("num2") or ""
    return unit.upper(), number


def is_dimension_connector(text: str) -> bool:
    return bool(re.fullmatch(r"\s*[xX*×]\s*", text or ""))


def connected_dimension_tokens(text: str, tokens: list[re.Match[str]], anchor_index: int) -> list[re.Match[str]]:
    start = anchor_index
    end = anchor_index

    while start > 0 and is_dimension_connector(text[tokens[start - 1].end() : tokens[start].start()]):
        start -= 1

    while end + 1 < len(tokens) and is_dimension_connector(text[tokens[end].end() : tokens[end + 1].start()]):
        end += 1

    return tokens[start : end + 1]


def split_dimension_description(text: str, start: int, end: int) -> tuple[str, str]:
    left = text[:start]
    right = text[end:]

    if right.lstrip().startswith(("x", "X", "*", "×")):
        suffix_start = len(right) - len(right.lstrip())
        return "", normalize_dimension_separators(right[suffix_start:])

    desc = clean_description(f"{left} {right}")
    return desc, ""


def build_ordered_spec(units: dict[str, str], desc: str = "", suffix: str = "") -> str:
    parts = []

    if "T" in units:
        parts.append(f"{format_number(units['T'])}T")
    if "H" in units:
        parts.append(f"{format_number(units['H'])}H")
    if "W" in units:
        parts.append(f"{format_number(units['W'])}W")
    if "L" in units:
        parts.append(f"{format_length_number(units['L'])}L")

    value = "x".join(parts) + suffix
    if desc:
        value = f"{value}_{desc}"
    return value


def format_trailing_dimension_token(token: str) -> str:
    paren_match = re.fullmatch(r"\((?P<body>[^)]*)\)", token.strip())
    if paren_match:
        body = normalize_dimension_separators(paren_match.group("body"))
        return f"({body})"

    match = re.fullmatch(rf"(?P<number>{NUMBER_PATTERN})(?P<suffix>[A-Za-z]+)?", token.strip(), re.IGNORECASE)
    if not match:
        return token.strip()
    return f"{format_optional_decimal_number(match.group('number'))}{match.group('suffix') or ''}"


def normalize_trailing_thickness_spec(original: str) -> NormalizeResult | None:
    text = normalize_dimension_separators(original)
    dimension_token = rf"(?:{NUMBER_PATTERN}(?:[A-Za-z]+)?|\([^)]*\))"
    match = re.search(rf"(?:(?<=x)|(?<![A-Za-z0-9-]))(?P<body>(?:{dimension_token}\s*x\s*)+)(?P<thick>{NUMBER_PATTERN})\s*T\b", text, re.IGNORECASE)
    if not match:
        return None

    dimension_tokens = [
        item.strip()
        for item in match.group("body").split("x")
        if item.strip()
    ]
    if not dimension_tokens:
        return None

    left = text[: match.start()].rstrip(" \t\r\n_x-/,:;")
    right = text[match.end() :].lstrip(" \t\r\n_x-/,:;")
    desc = clean_description(f"{left} {right}")
    parts = [f"{format_number(match.group('thick'))}T"]
    parts.extend(format_trailing_dimension_token(token) for token in dimension_tokens)
    value = "x".join(parts)
    if desc:
        value = f"{value}_{desc}"
    if value == original:
        return NormalizeResult("ok", value)
    return NormalizeResult("changed", value, "뒤쪽 T 두께를 앞으로 옮기고 단위 없는 숫자는 순서대로 유지했습니다.")


def normalize_spec_by_units(original: str) -> NormalizeResult | None:
    text = normalize_dimension_separators(original)
    tokens = list(UNIT_TOKEN_RE.finditer(text))
    t_indexes = [index for index, token in enumerate(tokens) if unit_token_value(token)[0] == "T"]
    h_indexes = [index for index, token in enumerate(tokens) if unit_token_value(token)[0] == "H"]

    if not t_indexes and not h_indexes:
        return None

    for anchor_index in [*t_indexes, *h_indexes]:
        component = connected_dimension_tokens(text, tokens, anchor_index)
        units: dict[str, str] = {}
        for token in component:
            unit, number = unit_token_value(token)
            units.setdefault(unit, number)

        if "T" not in units and not {"H", "W", "L"}.issubset(units):
            continue
        if set(units) == {"T"} and re.match(r"\s*[xX*×]\s*\d", text[component[-1].end() :]):
            continue

        desc, suffix = split_dimension_description(text, component[0].start(), component[-1].end())
        value = build_ordered_spec(units, desc=desc, suffix=suffix)
        if value == original:
            return NormalizeResult("ok", value)
        return NormalizeResult("changed", value, "T/H/W/L 단위를 정해진 순서로 정리했습니다.")

    return None


def clean_description(text: str) -> str:
    desc = re.sub(r"^[\s,;/_\-:]+|[\s,;/_\-:]+$", "", text)
    desc = re.sub(r"\s+", " ", desc)
    desc = re.sub(r"\s*_\s*", "_", desc)
    return desc


def normalize_category_text(value: Any) -> str:
    return cell_text(value).replace(" ", "").upper()


def should_skip_category(value: Any) -> bool:
    category = normalize_category_text(value)
    if category in SKIP_CATEGORY_EXACT:
        return True
    return any(category.startswith(prefix) for prefix in SKIP_CATEGORY_PREFIXES)


def remove_dimension_span(spec: str, start: int, end: int) -> str:
    unit_end = end
    while unit_end < len(spec) and spec[unit_end].isspace():
        unit_end += 1
    next_index = unit_end + 1
    if (
        unit_end < len(spec)
        and spec[unit_end].upper() in ("W", "L")
        and (next_index >= len(spec) or not spec[next_index].isalnum())
    ):
        end = unit_end + 1

    label_start = start
    while label_start > 0 and spec[label_start - 1].isspace():
        label_start -= 1
    if label_start > 0 and spec[label_start - 1].upper() in ("H", "T"):
        start = label_start - 1

    left = spec[:start]
    right = spec[end:]

    while left and left[-1] in " \t\r\n_xX*×-/,;:":
        left = left[:-1]
    while right and right[0] in " \t\r\n_xX*×-/,;:":
        right = right[1:]

    return clean_description(f"{left} {right}")


def remove_dimension_area(spec: str, matches: list[re.Match[str]]) -> str:
    return remove_dimension_span(spec, matches[0].start(), matches[-1].end())


def thickness_label(spec: str, first_number: re.Match[str]) -> str:
    before = spec[max(0, first_number.start() - 3) : first_number.start()].upper()
    after = spec[first_number.end() : first_number.end() + 2].upper()

    if "H" in before or "H" in after:
        return "H"
    return "T"


def normalize_spec(spec_value: Any) -> NormalizeResult:
    original = cell_text(spec_value)
    if not original:
        return NormalizeResult("skip", original, "규격이 비어 있습니다.")

    separator_normalized = normalize_dimension_separators(original)

    def manual_or_separator_change(reason: str) -> NormalizeResult:
        if separator_normalized != original:
            return NormalizeResult("changed", separator_normalized, "* 구분자를 x로 변경했습니다.")
        return NormalizeResult("manual", original, reason)

    trailing_thickness_result = normalize_trailing_thickness_spec(original)
    if trailing_thickness_result:
        return trailing_thickness_result

    unit_result = normalize_spec_by_units(original)
    if unit_result:
        return unit_result

    if not DIM_SEPARATOR_RE.search(original):
        return manual_or_separator_change("x 또는 * 구분자가 없는 규격입니다.")

    thickness_match = THICKNESS_SUFFIX_RE.search(original)
    if thickness_match is None:
        return manual_or_separator_change("두께 숫자 뒤에 T/H가 없어 자동 변경하지 않았습니다.")

    tail_start = thickness_match.end()
    tail_match = DIMENSION_TAIL_RE.match(original[tail_start:])
    if tail_match is None:
        return manual_or_separator_change("두께 뒤에 x폭 형식이 없어 자동 변경하지 않았습니다.")

    after = tail_match.group("after").lstrip()
    if after and DIM_SEPARATOR_RE.match(after):
        return manual_or_separator_change("길이 위치에 숫자가 없어 자동 변경하지 않았습니다.")

    length = tail_match.group("length")
    dimension_end = tail_start + (tail_match.end("length") if length else tail_match.end("width"))
    desc = remove_dimension_span(original, thickness_match.start(), dimension_end)

    width_unit = tail_match.group("width_unit")
    length_unit = tail_match.group("length_unit")
    if not width_unit and not length_unit:
        numbers = [tail_match.group("width")]
        if length:
            numbers.append(length)
        value = build_unlabeled_spec_value(
            thickness_match.group("label"),
            thickness_match.group("thick"),
            numbers,
            desc,
        )
    else:
        value = build_spec_value(
            thickness_match.group("label"),
            thickness_match.group("thick"),
            tail_match.group("width"),
            length,
            desc,
        )

    if value == original:
        if separator_normalized != original:
            return NormalizeResult("changed", separator_normalized, "* 구분자를 x로 변경했습니다.")
        return NormalizeResult("ok", value)
    if not width_unit and not length_unit:
        return NormalizeResult("changed", value, "단위 없는 숫자는 순서만 유지해 정리했습니다.")
    dimension_text = "두께/폭/길이" if length else "두께/폭"
    return NormalizeResult("changed", value, f"숫자와 단위를 {dimension_text}로 정리했습니다.")


def unique_sheet_name(workbook, base_name: str) -> str:
    if base_name not in workbook.sheetnames:
        return base_name

    for index in range(2, 1000):
        name = f"{base_name}_{index}"
        if name not in workbook.sheetnames:
            return name
    raise RuntimeError("결과 시트 이름을 만들 수 없습니다.")


def add_report_sheet(workbook, rows: list[dict[str, str]]) -> None:
    ws = workbook.create_sheet(unique_sheet_name(workbook, "규격정리결과"))
    headers = ["상태", "행", "품번", "기존 규격", "변경 규격", "사유"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row["status"],
                row["row"],
                row["part_number"],
                row["old_spec"],
                row["new_spec"],
                row["reason"],
            ]
        )

    for column_cells in ws.columns:
        max_len = max(len(cell_text(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 60)


def prepare_original_spec_column(ws, spec_col: int) -> int:
    original_col = spec_col + 1
    header_row = ERP_DATA_START_ROW - 1
    header_text = "원본 규격(확인용)"

    if cell_text(ws.cell(header_row, original_col).value) != header_text:
        ws.insert_cols(original_col)

    header_cell = ws.cell(header_row, original_col)
    header_cell.value = header_text
    header_cell.fill = HEADER_FILL
    header_cell.font = Font(bold=True)
    ws.column_dimensions[header_cell.column_letter].width = 24
    return original_col


def normalize_workbook_specs(
    input_file: str | Path,
    output_file: str | Path,
    db_path: str | Path = DEFAULT_ERP_DB,
    progress=None,
) -> dict[str, int]:
    part_keys = load_db_part_keys(db_path)

    if progress:
        progress(10, "엑셀 파일을 여는 중...")

    workbook = load_workbook(input_file)
    ws = workbook.worksheets[0]

    part_col = ERP_FIXED_COLUMNS["part_number"]
    spec_col = ERP_FIXED_COLUMNS["spec"]
    original_spec_col = prepare_original_spec_column(ws, spec_col)
    report_rows: list[dict[str, str]] = []
    matched = changed = manual = ok = skipped = 0
    max_row = ws.max_row

    for row_index in range(ERP_DATA_START_ROW, max_row + 1):
        if progress and row_index % 500 == 0:
            percent = 10 + int((row_index / max(max_row, 1)) * 75)
            progress(percent, f"규격 확인 중... {row_index:,}/{max_row:,}")

        part_cell = ws.cell(row_index, part_col)
        spec_cell = ws.cell(row_index, spec_col)
        part_number = cell_text(part_cell.value)
        part_key = normalize_part_lookup(part_number)

        if not part_key or part_key not in part_keys:
            continue

        matched += 1
        original_cell = ws.cell(row_index, original_spec_col)
        current_spec = cell_text(spec_cell.value)
        old_spec = current_spec
        result = normalize_spec(old_spec)

        if result.status == "skip":
            skipped += 1
            continue

        if result.status == "changed":
            spec_cell.value = result.value
            spec_cell.fill = CHANGED_FILL
            original_cell.value = old_spec
            original_cell.fill = CHANGED_FILL
            changed += 1
            report_rows.append(
                {
                    "status": "자동변경",
                    "row": str(row_index),
                    "part_number": part_number,
                    "old_spec": old_spec,
                    "new_spec": result.value,
                    "reason": result.reason,
                }
            )
        elif result.status == "manual":
            spec_cell.fill = MANUAL_FILL
            manual += 1
            report_rows.append(
                {
                    "status": "수동확인",
                    "row": str(row_index),
                    "part_number": part_number,
                    "old_spec": old_spec,
                    "new_spec": "",
                    "reason": result.reason,
                }
            )
        else:
            ok += 1

    if progress:
        progress(90, "결과 시트를 만드는 중...")

    add_report_sheet(workbook, report_rows)
    workbook.save(output_file)
    workbook.close()

    if progress:
        progress(100, "완료")

    return {
        "matched": matched,
        "changed": changed,
        "manual": manual,
        "ok": ok,
        "skipped": skipped,
        "reported": len(report_rows),
    }


def default_output_file(input_file: str | Path) -> str:
    source = Path(input_file)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(source.with_name(f"{source.stem}_규격정리_{timestamp}.xlsx"))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="ERP 품목 엑셀의 규격을 T/HxWxL 형식으로 정리하고 표시합니다.")
    parser.add_argument("excel_file", nargs="?", help="규격을 정리할 ERP 품목 엑셀 경로")
    parser.add_argument("-o", "--output", help="저장할 엑셀 경로")
    parser.add_argument("--erp-db", default=str(DEFAULT_ERP_DB), help="ERP 품목 DB 경로")
    return parser


def run_cli(excel_file: str, output_file: str | None, db_path: str) -> None:
    output = output_file or default_output_file(excel_file)
    stats = normalize_workbook_specs(excel_file, output, db_path)
    print(f"저장 완료: {output}")
    print(
        " / ".join(
            [
                f"DB 품번 매칭 {stats['matched']}",
                f"스킵 {stats['skipped']}",
                f"자동변경 {stats['changed']}",
                f"수동확인 {stats['manual']}",
                f"정상 {stats['ok']}",
            ]
        )
    )


def run_gui() -> None:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        initial_file = latest_db_source_file(DEFAULT_ERP_DB)
        excel_file = select_excel_file("규격을 정리할 ERP 품목 엑셀 선택", parent=root, initial_path=initial_file)
        if not excel_file:
            return

        output_file = select_excel_save_file("규격 정리 결과 저장", excel_file, "규격정리", parent=root)
        if not output_file:
            return

        progress_win, update_progress = create_progress_window(root, "ERP 규격 정리")
        try:
            stats = normalize_workbook_specs(excel_file, output_file, DEFAULT_ERP_DB, progress=update_progress)
        finally:
            progress_win.destroy()

        messagebox.showinfo(
            "완료",
            f"규격 정리 파일을 저장했습니다.\n\n"
            f"파일: {output_file}\n"
            f"DB 품번 매칭: {stats['matched']}\n"
            f"스킵: {stats['skipped']}\n"
            f"자동 변경(노랑): {stats['changed']}\n"
            f"수동 확인(빨강): {stats['manual']}\n"
            f"이미 정상: {stats['ok']}",
            parent=root,
        )
    except Exception as exc:
        messagebox.showerror("오류", str(exc), parent=root)
    finally:
        root.destroy()


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.excel_file:
        run_cli(args.excel_file, args.output, args.erp_db)
        return

    run_gui()


if __name__ == "__main__":
    main()
