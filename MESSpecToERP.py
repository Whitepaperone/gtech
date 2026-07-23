from __future__ import annotations

import argparse
import tkinter as tk
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tkinter import messagebox
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from ChangeListParser import cell_text
from CommonUtils import create_progress_window, select_excel_file, select_excel_save_file
from ERPPartsDb import ERP_DATA_START_ROW, ERP_FIXED_COLUMNS, normalize_part_lookup


MES_DATA_START_ROW = 2
MES_PART_COL = 1  # A
MES_SPEC_COL = 4  # D
ERP_PART_COL = ERP_FIXED_COLUMNS["part_number"]  # B
ERP_SPEC_COL = ERP_FIXED_COLUMNS["spec"]  # E

CHANGED_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")

try:
    import pythoncom
    import win32com.client
except ImportError:  # pragma: no cover - fallback for PCs without Excel COM support.
    pythoncom = None
    win32com = None


@dataclass(frozen=True)
class ErpSpec:
    part_number: str
    spec: str
    row: int


def unique_sheet_name(workbook, base_name: str) -> str:
    if base_name not in workbook.sheetnames:
        return base_name

    for index in range(2, 1000):
        name = f"{base_name}_{index}"
        if name not in workbook.sheetnames:
            return name
    raise RuntimeError("Could not create a unique result sheet name.")


def default_output_file(mes_file: str | Path) -> str:
    source = Path(mes_file)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return str(source.with_name(f"{source.stem}_ERP_spec_applied_{timestamp}.xlsx"))


def normalize_com_range_values(values: Any) -> list[tuple[Any, ...]]:
    if values is None:
        return []
    if not isinstance(values, tuple):
        return [(values,)]
    if not values:
        return []
    if not isinstance(values[0], tuple):
        return [values]
    return list(values)


def add_erp_spec(
    specs: dict[str, ErpSpec],
    duplicate_rows: list[dict[str, str]],
    part_number: Any,
    spec_value: Any,
    row_index: int,
) -> bool:
    part_number_text = cell_text(part_number)
    spec = cell_text(spec_value)
    part_key = normalize_part_lookup(part_number_text)

    if not part_key or not spec:
        return False

    previous = specs.get(part_key)
    if previous is None:
        specs[part_key] = ErpSpec(part_number=part_number_text, spec=spec, row=row_index)
        return True

    if previous.spec != spec:
        duplicate_rows.append(
            {
                "part_number": part_number_text,
                "first_row": str(previous.row),
                "duplicate_row": str(row_index),
                "kept_spec": previous.spec,
                "ignored_spec": spec,
            }
        )
    return True


def raise_no_erp_specs(stats_text: str) -> None:
    raise ValueError(
        "ERP file has no rows where both part code column B and spec column E are filled.\n\n"
        f"Scan result:\n{stats_text}\n\n"
        "Please check whether the ERP part/spec columns are really B/E, or whether spec cells are blank."
    )


def load_erp_specs_with_excel_com(erp_file: str | Path, progress=None) -> tuple[dict[str, ErpSpec], list[dict[str, str]]]:
    if pythoncom is None or win32com is None:
        raise RuntimeError("Excel COM is not available.")

    source_file = str(Path(erp_file).resolve())
    specs: dict[str, ErpSpec] = {}
    duplicate_rows: list[dict[str, str]] = []

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(source_file, ReadOnly=True)
        ws = workbook.Worksheets(1)
        last_part_row = ws.Cells(ws.Rows.Count, ERP_PART_COL).End(-4162).Row
        last_spec_row = ws.Cells(ws.Rows.Count, ERP_SPEC_COL).End(-4162).Row
        last_row = max(last_part_row, last_spec_row)
        part_count = spec_count = pair_count = 0

        if progress:
            progress(15, f"Fast ERP scan... {last_row:,} rows")

        if last_row >= ERP_DATA_START_ROW:
            rows = normalize_com_range_values(ws.Range(f"B{ERP_DATA_START_ROW}:E{last_row}").Value)
            for offset, row_values in enumerate(rows):
                row_index = ERP_DATA_START_ROW + offset
                part_number = row_values[0] if len(row_values) > 0 else ""
                spec_value = row_values[3] if len(row_values) > 3 else ""
                part_key = normalize_part_lookup(part_number)
                spec = cell_text(spec_value)

                if part_key:
                    part_count += 1
                if spec:
                    spec_count += 1
                if add_erp_spec(specs, duplicate_rows, part_number, spec_value, row_index):
                    pair_count += 1
    finally:
        if workbook is not None:
            workbook.Close(False)
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()

    if not specs:
        raise_no_erp_specs(f"ERP first sheet: B={part_count}, E={spec_count}, B+E={pair_count}")

    return specs, duplicate_rows


def load_erp_specs_with_openpyxl(erp_file: str | Path, progress=None) -> tuple[dict[str, ErpSpec], list[dict[str, str]]]:
    workbook = load_workbook(erp_file, read_only=True, data_only=True)
    specs: dict[str, ErpSpec] = {}
    duplicate_rows: list[dict[str, str]] = []
    part_count = spec_count = pair_count = 0

    try:
        ws = workbook.worksheets[0]
        max_row = ws.max_row or 1
        if progress:
            progress(15, f"ERP scan... {ws.title}")

        for row_index in range(ERP_DATA_START_ROW, max_row + 1):
            if progress and row_index % 5000 == 0:
                progress(15, f"ERP scan... {row_index:,}/{max_row:,}")

            part_number = ws.cell(row_index, ERP_PART_COL).value
            spec_value = ws.cell(row_index, ERP_SPEC_COL).value
            part_key = normalize_part_lookup(part_number)
            spec = cell_text(spec_value)

            if part_key:
                part_count += 1
            if spec:
                spec_count += 1
            if add_erp_spec(specs, duplicate_rows, part_number, spec_value, row_index):
                pair_count += 1
    finally:
        workbook.close()

    if not specs:
        raise_no_erp_specs(f"ERP first sheet: B={part_count}, E={spec_count}, B+E={pair_count}")

    return specs, duplicate_rows


def load_erp_specs(erp_file: str | Path, progress=None) -> tuple[dict[str, ErpSpec], list[dict[str, str]]]:
    try:
        return load_erp_specs_with_excel_com(erp_file, progress=progress)
    except Exception:
        return load_erp_specs_with_openpyxl(erp_file, progress=progress)


def add_report_sheet(
    workbook,
    changed_rows: list[dict[str, str]],
    unmatched_rows: list[dict[str, str]],
    duplicate_rows: list[dict[str, str]],
) -> None:
    ws = workbook.create_sheet(unique_sheet_name(workbook, "ERP_spec_result"))
    headers = ["Status", "MES sheet", "MES row", "Part number", "Old MES spec", "ERP spec", "ERP row", "Note"]
    ws.append(headers)

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)

    for row in changed_rows:
        ws.append(
            [
                "Changed",
                row["mes_sheet"],
                row["mes_row"],
                row["part_number"],
                row["old_spec"],
                row["new_spec"],
                row["erp_row"],
                "",
            ]
        )

    for row in unmatched_rows:
        ws.append(
            [
                "No ERP match",
                row["mes_sheet"],
                row["mes_row"],
                row["part_number"],
                row["old_spec"],
                "",
                "",
                "No ERP part code matched this MES part number.",
            ]
        )

    for row in duplicate_rows:
        ws.append(
            [
                "ERP duplicate",
                "",
                "",
                row["part_number"],
                row["kept_spec"],
                row["ignored_spec"],
                row["duplicate_row"],
                f"First ERP spec from row {row['first_row']} was used.",
            ]
        )

    for column_cells in ws.columns:
        max_len = max(len(cell_text(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 70)


def excel_file_format(output_file: str | Path) -> int:
    suffix = Path(output_file).suffix.lower()
    if suffix == ".xlsm":
        return 52
    if suffix == ".xls":
        return 56
    return 51


def com_unique_sheet_name(workbook, base_name: str) -> str:
    existing = {str(workbook.Worksheets(index).Name) for index in range(1, workbook.Worksheets.Count + 1)}
    if base_name not in existing:
        return base_name

    for index in range(2, 1000):
        name = f"{base_name}_{index}"
        if name not in existing:
            return name
    raise RuntimeError("Could not create a unique result sheet name.")


def add_report_sheet_with_excel_com(
    workbook,
    changed_rows: list[dict[str, str]],
    unmatched_rows: list[dict[str, str]],
    duplicate_rows: list[dict[str, str]],
) -> None:
    ws = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
    ws.Name = com_unique_sheet_name(workbook, "ERP_spec_result")

    rows = [["Status", "MES sheet", "MES row", "Part number", "Old MES spec", "ERP spec", "ERP row", "Note"]]
    for row in changed_rows:
        rows.append(
            [
                "Changed",
                row["mes_sheet"],
                row["mes_row"],
                row["part_number"],
                row["old_spec"],
                row["new_spec"],
                row["erp_row"],
                "",
            ]
        )

    for row in unmatched_rows:
        rows.append(
            [
                "No ERP match",
                row["mes_sheet"],
                row["mes_row"],
                row["part_number"],
                row["old_spec"],
                "",
                "",
                "No ERP part code matched this MES part number.",
            ]
        )

    for row in duplicate_rows:
        rows.append(
            [
                "ERP duplicate",
                "",
                "",
                row["part_number"],
                row["kept_spec"],
                row["ignored_spec"],
                row["duplicate_row"],
                f"First ERP spec from row {row['first_row']} was used.",
            ]
        )

    if rows:
        ws.Range(ws.Cells(1, 1), ws.Cells(len(rows), len(rows[0]))).Value = tuple(tuple(row) for row in rows)
        ws.Range(ws.Cells(1, 1), ws.Cells(1, len(rows[0]))).Font.Bold = True
        ws.Range(ws.Cells(1, 1), ws.Cells(1, len(rows[0]))).Interior.Color = 16247773
        ws.Columns.AutoFit()


def replace_mes_specs_from_erp_with_excel_com(
    mes_file: str | Path,
    erp_file: str | Path,
    output_file: str | Path,
    progress=None,
) -> dict[str, int]:
    if pythoncom is None or win32com is None:
        raise RuntimeError("Excel COM is not available.")

    if progress:
        progress(5, "Reading ERP product file...")
    erp_specs, duplicate_rows = load_erp_specs(erp_file, progress=progress)

    source_file = str(Path(mes_file).resolve())
    output_path = str(Path(output_file).resolve())
    changed_rows: list[dict[str, str]] = []
    unmatched_rows: list[dict[str, str]] = []
    matched = changed = same = missing = blank_part = 0

    pythoncom.CoInitialize()
    excel = None
    workbook = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False
        excel.EnableEvents = False
        workbook = excel.Workbooks.Open(source_file)
        sheet_count = max(workbook.Worksheets.Count, 1)

        for sheet_index in range(1, sheet_count + 1):
            ws = workbook.Worksheets(sheet_index)
            sheet_name = str(ws.Name)
            last_part_row = ws.Cells(ws.Rows.Count, MES_PART_COL).End(-4162).Row
            last_spec_row = ws.Cells(ws.Rows.Count, MES_SPEC_COL).End(-4162).Row
            max_row = max(last_part_row, last_spec_row)
            if max_row < MES_DATA_START_ROW:
                continue

            if progress:
                percent = 35 + int((sheet_index - 1) / sheet_count * 50)
                progress(percent, f"Fast MES update... {sheet_name}")

            rows = normalize_com_range_values(ws.Range(f"A{MES_DATA_START_ROW}:D{max_row}").Value)
            for offset, row_values in enumerate(rows):
                row_index = MES_DATA_START_ROW + offset
                part_number = cell_text(row_values[0] if len(row_values) > 0 else "")
                part_key = normalize_part_lookup(part_number)

                if not part_key:
                    blank_part += 1
                    continue

                erp_spec = erp_specs.get(part_key)
                old_spec = cell_text(row_values[3] if len(row_values) > 3 else "")
                if erp_spec is None:
                    missing += 1
                    unmatched_rows.append(
                        {
                            "mes_sheet": sheet_name,
                            "mes_row": str(row_index),
                            "part_number": part_number,
                            "old_spec": old_spec,
                        }
                    )
                    continue

                matched += 1
                if old_spec == erp_spec.spec:
                    same += 1
                    continue

                ws.Cells(row_index, MES_SPEC_COL).Value = erp_spec.spec
                ws.Cells(row_index, MES_SPEC_COL).Interior.Color = 13431551
                changed += 1
                changed_rows.append(
                    {
                        "mes_sheet": sheet_name,
                        "mes_row": str(row_index),
                        "part_number": part_number,
                        "old_spec": old_spec,
                        "new_spec": erp_spec.spec,
                        "erp_row": str(erp_spec.row),
                    }
                )

        if progress:
            progress(88, "Creating result sheet...")
        add_report_sheet_with_excel_com(workbook, changed_rows, unmatched_rows, duplicate_rows)

        if progress:
            progress(95, "Saving with Excel...")
        workbook.SaveAs(output_path, FileFormat=excel_file_format(output_path))
    finally:
        if workbook is not None:
            workbook.Close(False)
        if excel is not None:
            excel.EnableEvents = True
            excel.ScreenUpdating = True
            excel.Quit()
        pythoncom.CoUninitialize()

    if progress:
        progress(100, "Done")

    return {
        "erp_specs": len(erp_specs),
        "matched": matched,
        "changed": changed,
        "same": same,
        "missing": missing,
        "blank_part": blank_part,
        "duplicate_conflicts": len(duplicate_rows),
    }


def replace_mes_specs_from_erp_with_openpyxl(
    mes_file: str | Path,
    erp_file: str | Path,
    output_file: str | Path,
    progress=None,
) -> dict[str, int]:
    if progress:
        progress(5, "Reading ERP product file...")
    erp_specs, duplicate_rows = load_erp_specs(erp_file, progress=progress)

    if progress:
        progress(35, "Opening MES file...")

    mes_path = Path(mes_file)
    keep_vba = mes_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(mes_file, keep_vba=keep_vba)

    changed_rows: list[dict[str, str]] = []
    unmatched_rows: list[dict[str, str]] = []
    matched = changed = same = missing = blank_part = 0

    try:
        sheet_count = max(len(workbook.worksheets), 1)
        for sheet_index, ws in enumerate(workbook.worksheets, start=1):
            max_row = ws.max_row or 1
            for row_index in range(MES_DATA_START_ROW, max_row + 1):
                if progress and row_index % 1000 == 0:
                    sheet_base = 35 + int((sheet_index - 1) / sheet_count * 55)
                    sheet_span = 55 / sheet_count
                    percent = min(90, sheet_base + int((row_index / max(max_row, 1)) * sheet_span))
                    progress(percent, f"Applying ERP specs to MES... {ws.title} {row_index:,}/{max_row:,}")

                part_cell = ws.cell(row_index, MES_PART_COL)
                spec_cell = ws.cell(row_index, MES_SPEC_COL)
                part_number = cell_text(part_cell.value)
                part_key = normalize_part_lookup(part_number)

                if not part_key:
                    blank_part += 1
                    continue

                erp_spec = erp_specs.get(part_key)
                old_spec = cell_text(spec_cell.value)
                if erp_spec is None:
                    missing += 1
                    unmatched_rows.append(
                        {
                            "mes_sheet": ws.title,
                            "mes_row": str(row_index),
                            "part_number": part_number,
                            "old_spec": old_spec,
                        }
                    )
                    continue

                matched += 1
                if old_spec == erp_spec.spec:
                    same += 1
                    continue

                spec_cell.value = erp_spec.spec
                spec_cell.fill = CHANGED_FILL
                changed += 1
                changed_rows.append(
                    {
                        "mes_sheet": ws.title,
                        "mes_row": str(row_index),
                        "part_number": part_number,
                        "old_spec": old_spec,
                        "new_spec": erp_spec.spec,
                        "erp_row": str(erp_spec.row),
                    }
                )

        if progress:
            progress(92, "Creating result sheet...")

        add_report_sheet(workbook, changed_rows, unmatched_rows, duplicate_rows)
        workbook.save(output_file)
    finally:
        workbook.close()

    if progress:
        progress(100, "Done")

    return {
        "erp_specs": len(erp_specs),
        "matched": matched,
        "changed": changed,
        "same": same,
        "missing": missing,
        "blank_part": blank_part,
        "duplicate_conflicts": len(duplicate_rows),
    }


def replace_mes_specs_from_erp(
    mes_file: str | Path,
    erp_file: str | Path,
    output_file: str | Path,
    progress=None,
) -> dict[str, int]:
    try:
        return replace_mes_specs_from_erp_with_excel_com(mes_file, erp_file, output_file, progress=progress)
    except Exception:
        return replace_mes_specs_from_erp_with_openpyxl(mes_file, erp_file, output_file, progress=progress)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Apply ERP spec column E to MES spec column D by part number.")
    parser.add_argument("mes_file", nargs="?", help="MES product Excel file path")
    parser.add_argument("erp_file", nargs="?", help="ERP product Excel file path")
    parser.add_argument("-o", "--output", help="Output MES Excel file path")
    return parser


def run_cli(mes_file: str, erp_file: str, output_file: str | None) -> None:
    output = output_file or default_output_file(mes_file)
    stats = replace_mes_specs_from_erp(mes_file, erp_file, output)
    print(f"Saved: {output}")
    print(
        " / ".join(
            [
                f"ERP specs {stats['erp_specs']}",
                f"MES matched {stats['matched']}",
                f"Changed {stats['changed']}",
                f"Same {stats['same']}",
                f"No ERP match {stats['missing']}",
                f"ERP duplicate conflicts {stats['duplicate_conflicts']}",
            ]
        )
    )


def run_gui() -> None:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    try:
        mes_file = select_excel_file("Select MES product file to update", parent=root)
        if not mes_file:
            return

        erp_file = select_excel_file("Select ERP product file with specs", parent=root)
        if not erp_file:
            return

        output_file = select_excel_save_file("Save MES file with ERP specs", mes_file, "ERP_spec_applied", parent=root)
        if not output_file:
            return

        progress_win, update_progress = create_progress_window(root, "Apply ERP specs to MES")
        try:
            stats = replace_mes_specs_from_erp(mes_file, erp_file, output_file, progress=update_progress)
        finally:
            progress_win.destroy()

        messagebox.showinfo(
            "Done",
            f"ERP specs were applied to the MES file.\n\n"
            f"File: {output_file}\n"
            f"ERP specs: {stats['erp_specs']}\n"
            f"MES matched: {stats['matched']}\n"
            f"Changed: {stats['changed']}\n"
            f"Same: {stats['same']}\n"
            f"No ERP match: {stats['missing']}\n"
            f"ERP duplicate conflicts: {stats['duplicate_conflicts']}",
            parent=root,
        )
    except Exception as exc:
        messagebox.showerror("Error", str(exc), parent=root)
        raise
    finally:
        root.destroy()


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.mes_file and args.erp_file:
        run_cli(args.mes_file, args.erp_file, args.output)
        return

    run_gui()


if __name__ == "__main__":
    main()
